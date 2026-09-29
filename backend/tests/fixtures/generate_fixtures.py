from __future__ import annotations

import csv
import json
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

FIXTURE_DIR = Path(__file__).parent


def _make_text_pdf() -> bytes:
    lines = [
        "Project: Document Checker Sample.",
        "Purpose: validate text extraction and source page links.",
        "Author: Alex Example. Reviewer: Morgan Example.",
        "Technology: Python, FastAPI, and PostgreSQL.",
        "Decision: keep the local archive for seven years.",
        "Due date: 2026-11-30.",
    ]
    commands = ["BT", "/F1 12 Tf", "72 720 Td"]
    for index, line in enumerate(lines):
        if index:
            commands.append("0 -18 Td")
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"({escaped}) Tj")
    commands.append("ET")
    content = "\n".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n".encode("ascii") + content + b"\nendstream",
    ]

    document = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, value in enumerate(objects, start=1):
        offsets.append(len(document))
        document.extend(f"{number} 0 obj\n".encode("ascii"))
        document.extend(value)
        document.extend(b"\nendobj\n")

    xref_offset = len(document)
    document.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    document.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        document.extend(f"{offset:010} 00000 n \n".encode("ascii"))
    document.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(document)


def _make_docx() -> bytes:
    document = Document()
    document.add_heading("Тестовый проект", level=1)
    document.add_paragraph("Цель: проверить обработку документа и ссылки на источники.")
    document.add_paragraph("Автор: Алексей Пример. Проверяющий: Мария Пример.")
    document.add_paragraph("Технологии: Python, FastAPI и PostgreSQL.")
    document.add_paragraph("Срок: 30 ноября 2026 года.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Этап"
    table.cell(0, 1).text = "Ответственный"
    cells = table.add_row().cells
    cells[0].text = "Проверка"
    cells[1].text = "Алексей Пример"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _make_csv() -> bytes:
    stream = StringIO(newline="")
    writer = csv.writer(stream, delimiter=";", lineterminator="\n")
    writer.writerows([
        ["Проект", "Часы", "Стоимость"],
        ["Альфа", "10", "125.25"],
        ["Бета", "20", "80.25"],
        ["Гамма", "30", "94.50"],
    ])
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _make_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Данные"
    sheet.append(["Проект", "Часы", "Статус"])
    sheet.append(["Альфа", 10, "Готово"])
    sheet.append(["Бета", 20, "В работе"])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _make_xls() -> bytes:
    import xlwt
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Данные")
    for column, value in enumerate(["Проект", "Часы", "Статус"]):
        sheet.write(0, column, value)
    for row, values in enumerate([["Альфа", 10, "Готово"], ["Бета", 20, "В работе"]], start=1):
        for column, value in enumerate(values):
            sheet.write(row, column, value)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _make_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1))
    title.text = "Тестовый проект"
    body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(2))
    body.text = "Цель: проверить обработку презентации.\nСрок: 30 ноября 2026 года."
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def _make_epub() -> bytes:
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/epub+zip", compress_type=0)
        archive.writestr("META-INF/container.xml", "<?xml version=\"1.0\"?><container version=\"1.0\" xmlns=\"urn:oasis:names:tc:opendocument:xmlns:container\"><rootfiles><rootfile full-path=\"OEBPS/content.opf\" media-type=\"application/oebps-package+xml\"/></rootfiles></container>")
        archive.writestr("OEBPS/content.opf", "<?xml version=\"1.0\"?><package xmlns=\"http://www.idpf.org/2007/opf\" version=\"3.0\"><metadata xmlns:dc=\"http://purl.org/dc/elements/1.1/\"><dc:title>Тестовый проект</dc:title><dc:language>ru</dc:language></metadata><manifest><item id=\"chapter\" href=\"chapter1.xhtml\" media-type=\"application/xhtml+xml\"/></manifest><spine><itemref idref=\"chapter\"/></spine></package>")
        archive.writestr("OEBPS/chapter1.xhtml", "<html><body><h1>Тестовый проект</h1><p>Цель: проверить EPUB и ссылки на главу.</p></body></html>")
    return stream.getvalue()


FIXTURE_CONTENTS = {
    "sample.pdf": _make_text_pdf(),
    "sample.docx": _make_docx(),
    "sample.txt": (
        "Проект: Проверка документов\n"
        "Цель: найти автора и сроки в тексте.\n"
        "Автор: Алексей Пример; проверяющий: Мария Пример.\n"
        "Технологии: Python и PostgreSQL.\n"
        "Срок: 30 ноября 2026 года.\n"
    ).encode(),
    "sample.md": (
        "# Тестовый проект\n\n"
        "## Цель\nПроверить извлечение Markdown и диапазоны строк.\n\n"
        "## Участники\nАвтор — Алексей Пример; исполнитель — команда разработки.\n\n"
        "## Срок\n30 ноября 2026 года.\n"
    ).encode(),
    "sample.csv": _make_csv(),
    "sample.xlsx": _make_xlsx(),
    "sample.xls": _make_xls(),
    "sample.pptx": _make_pptx(),
    "sample.html": "<html><body><h1>Тестовый проект</h1><p>Цель: проверить HTML.</p></body></html>".encode(),
    "sample.json": json.dumps({"project": "Тестовый проект", "owner": "Алексей Пример", "due": "2026-11-30"}, ensure_ascii=False).encode(),
    "sample.epub": _make_epub(),
    "sample.xml": (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<project id="DEMO-001">\n'
        "  <title>Тестовый проект</title>\n"
        '  <owner role="executor">Алексей Пример</owner>\n'
        "  <technology>Python и PostgreSQL</technology>\n"
        '  <milestone due="2026-11-30">Проверка</milestone>\n'
        "</project>\n"
    ).encode(),
}


def main() -> None:
    for name, content in FIXTURE_CONTENTS.items():
        (FIXTURE_DIR / name).write_bytes(content)
        print(f"{name}: {len(content):,} bytes")


if __name__ == "__main__":
    main()
