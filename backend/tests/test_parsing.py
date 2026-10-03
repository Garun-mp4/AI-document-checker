from __future__ import annotations

import csv
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.markdown_mapping import add_xlsx_formula_context, map_markdown
from app.services.markitdown_service import MarkItDownService
from app.services.parsing import DocumentParsingError, parse_document, safe_filename


def _xlsx_with_formula_states() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Данные"
    sheet.append(["Позиция", "Количество", "Итог"])
    sheet.append(["A", 2, "=B2*10"])
    sheet.append([None, None, None])
    sheet.append(["B", 3, "=B4*10"])
    sheet.append(["Literal NaN", "NaN", None])
    sheet.append(["Zero result", 0, "=B6-B6"])
    totals = workbook.create_sheet("Итоги")
    totals.append(["Показатель", "Значение"])
    totals.append(["Проверка", "=2+2"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    cached = BytesIO()
    with zipfile.ZipFile(BytesIO(stream.getvalue())) as source, zipfile.ZipFile(cached, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<f>B2*10</f><v></v>", b"<f>B2*10</f><v>20</v>")
                content = content.replace(b"<f>B6-B6</f><v></v>", b"<f>B6-B6</f><v>0</v>")
            target.writestr(item, content)
    return cached.getvalue()


def test_safe_filename_keeps_only_leaf_name() -> None:
    assert safe_filename(r"C:\private\folder\report.pdf") == "report.pdf"


def test_plain_text_supports_cp1251_and_keeps_line_locations() -> None:
    parsed = parse_document("notes.txt", "Первая строка\nВторая строка".encode("cp1251"))

    assert parsed.file_type == "txt"
    assert "Первая строка" in parsed.blocks[0].text
    assert parsed.blocks[0].locator["line_start"] == 1
    assert parsed.blocks[0].locator["line_end"] == 2


def test_plain_text_crlf_ranges_point_to_exact_original_characters() -> None:
    original = "Заголовок\r\nЦитируемая строка документа\r\nЗаключение"
    parsed = parse_document("ranges.txt", original.encode("utf-8"))
    source = next(block for block in parsed.blocks if "Цитируемая строка" in block.text)

    excerpt = original[source.locator["char_start"]:source.locator["char_end"]]
    assert "Цитируемая строка" in excerpt
    assert source.locator["line_start"] == 1
    assert source.locator["line_end"] == 3


def test_csv_fallback_does_not_mutate_global_excel_dialect(monkeypatch: pytest.MonkeyPatch) -> None:
    def no_dialect(*args: object, **kwargs: object) -> csv.Dialect:
        raise csv.Error("ambiguous sample")

    monkeypatch.setattr(csv.Sniffer, "sniff", no_dialect)
    parsed = parse_document("table.csv", b"Name;Value\none;first\ntwo;second\n")

    assert parsed.metadata["delimiter"] == ";"
    assert csv.excel.delimiter == ","
    assert any("Name: one" in block.text for block in parsed.blocks)


def test_csv_calculates_decimal_comma_columns_exactly() -> None:
    parsed = parse_document("table.csv", "Товар;Сумма\nА;1,5\nБ;2,5\n".encode())

    assert parsed.metadata["row_count"] == 2
    assert parsed.metadata["column_count"] == 2
    assert parsed.metadata["numeric_columns"] == [{
        "name": "Сумма",
        "count": 2,
        "sum": "4.0",
        "average": "2.0",
        "minimum": "1.5",
        "maximum": "2.5",
    }]


def test_xlsx_parser_keeps_formula_and_cached_value_states_with_cell_locators() -> None:
    parsed = parse_document("formulas.xlsx", _xlsx_with_formula_states())
    by_cell = {
        formula["cell"]: (block, formula)
        for block in parsed.blocks
        for formula in block.locator.get("formula_cells", [])
    }

    cached_block, cached_formula = by_cell["C2"]
    assert cached_formula == {
        "cell": "C2",
        "column": "Итог",
        "column_index": 2,
        "formula": "=B2*10",
        "has_cached_value": True,
        "cached_value": "20",
    }
    assert "Количество: 2" in cached_block.text
    assert "формула: =B2*10" in cached_block.text
    assert "сохранённое значение: 20" in cached_block.text

    uncached_block, uncached_formula = by_cell["C4"]
    assert uncached_formula["has_cached_value"] is False
    assert uncached_formula["cached_value"] is None
    assert "формула: =B4*10" in uncached_block.text
    assert "сохранённое значение отсутствует" in uncached_block.text
    assert "не вычисля" in uncached_block.text
    assert uncached_block.locator["sheet"] == "Данные"
    assert uncached_block.locator["row_start"] == 4
    assert uncached_block.locator["columns"] == ["Позиция", "Количество", "Итог"]

    second_sheet_block, second_sheet_formula = by_cell["B2"]
    assert second_sheet_block.locator["sheet"] == "Итоги"
    assert second_sheet_formula["formula"] == "=2+2"
    zero_formula = next(
        formula for block in parsed.blocks for formula in block.locator.get("formula_cells", [])
        if formula["cell"] == "C6"
    )
    assert zero_formula["has_cached_value"] is True
    assert zero_formula["cached_value"] == "0"
    assert parsed.metadata["formula_count"] == 4
    assert parsed.metadata["formula_cache_missing_count"] == 2


def test_xlsx_parser_extracts_array_formula_expression_without_evaluating_it() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Input", "Result"])
    sheet.append([2, None])
    sheet["B2"] = ArrayFormula(ref="B2:B2", text="=A2*2")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    parsed = parse_document("array-formula.xlsx", stream.getvalue())
    formula = next(
        formula for block in parsed.blocks for formula in block.locator.get("formula_cells", [])
    )

    assert formula["formula"] == "=A2*2"
    assert formula["has_cached_value"] is False
    assert parsed.metadata["formula_cache_missing_count"] == 1
    assert "приложение формулу не вычисляло" in next(block.text for block in parsed.blocks if formula in block.locator.get("formula_cells", []))


def test_xlsx_markdown_reconciles_empty_cells_and_maps_formula_citations(tmp_path: Path) -> None:
    content = _xlsx_with_formula_states()
    source = tmp_path / "formulas.xlsx"
    source.write_bytes(content)
    native = parse_document(source.name, content)
    converted = MarkItDownService(tmp_path).convert_local(source).markdown

    markdown = add_xlsx_formula_context(converted, native.blocks)

    assert "| A | 2.0 | 20 |" in markdown
    assert "|  |  |  |" in markdown
    assert "| B | 3.0 | Нет сохранённого результата |" in markdown
    assert "| Literal NaN | NaN |  |" in markdown
    assert "выражение `=B2*10`" in markdown
    assert "Формула XLSX — лист «Данные», ячейка C6" in markdown
    assert "сохранённое значение: 0" in markdown
    assert "сохранённое значение отсутствует; приложение формулу не вычисляло" in markdown
    mapped, _mapping = map_markdown(markdown, native.blocks)
    cached_citation = next(block for block in mapped if block.locator.get("cell") == "C2")
    uncached_citation = next(block for block in mapped if block.locator.get("cell") == "C4")
    zero_citation = next(block for block in mapped if block.locator.get("cell") == "C6")
    assert cached_citation.locator["sheet"] == "Данные"
    assert cached_citation.locator["row_start"] == 2
    assert cached_citation.locator["column_index"] == 2
    assert cached_citation.locator["mapping_confidence"] == "exact"
    assert uncached_citation.locator["row_start"] == 4
    assert uncached_citation.locator["column_index"] == 2
    assert uncached_citation.locator["formula_has_cached_value"] is False
    assert zero_citation.locator["row_start"] == 6
    assert zero_citation.locator["formula_has_cached_value"] is True


def test_large_csv_row_is_split_without_losing_its_locator() -> None:
    data = ("Описание\n" + "слово " * 400).encode()
    parsed = parse_document("large.csv", data)
    row_parts = [block for block in parsed.blocks if block.locator.get("row_start") == 2]

    assert len(row_parts) > 1
    assert all(block.locator["row_end"] == 2 for block in row_parts)
    assert all(len(block.text) <= 900 for block in row_parts)


def test_docx_preserves_paragraph_and_table_row_locations() -> None:
    document = DocxDocument()
    document.add_heading("Цель", level=1)
    document.add_paragraph("Проверить документ.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Этап"
    table.cell(0, 1).text = "Срок"
    stream = BytesIO()
    document.save(stream)
    parsed = parse_document("brief.docx", stream.getvalue())

    assert any(block.locator.get("heading") == "Heading 1" for block in parsed.blocks)
    assert any(block.locator.get("kind") == "docx_table" and "Срок" in block.text for block in parsed.blocks)


def test_xml_keeps_element_path_and_rejects_entities() -> None:
    original = "<report><owner>Team</owner></report>"
    parsed = parse_document("report.xml", original.encode())

    assert parsed.blocks[0].locator["path"] == "/report[1]/owner[1]"
    assert "Team" in original[parsed.blocks[0].locator["char_start"]:parsed.blocks[0].locator["char_end"]]
    with pytest.raises(DocumentParsingError, match="XML"):
        parse_document("unsafe.xml", b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///secret">]><x>&e;</x>')


def test_xml_entity_ranges_cover_the_original_encoded_value() -> None:
    original = "<report>\n  <owner>Team &amp; Alpha</owner>\n</report>"
    parsed = parse_document("entity-value.xml", original.encode())
    source = next(block for block in parsed.blocks if "Team & Alpha" in block.text)

    value = original[source.locator["char_start"]:source.locator["char_end"]]
    assert "Team &amp; Alpha" in value
    assert source.locator["line_start"] == 2


def test_json_locators_include_exact_raw_value_ranges() -> None:
    original = '{\n  "owner": "Алексей Пример",\n  "items": [1, 2]\n}'
    parsed = parse_document("report.json", original.encode())
    source = next(block for block in parsed.blocks if block.locator.get("path") == "$.owner")

    value = original[source.locator["char_start"]:source.locator["char_end"]]
    assert "Алексей Пример" in value
    assert source.locator["line_start"] == source.locator["line_end"] == 2


def test_html_locators_include_original_source_ranges() -> None:
    original = "<html>\n<body>\n<p>Автор: Алексей &amp; команда</p>\n</body>\n</html>"
    parsed = parse_document("report.html", original.encode())
    source = next(block for block in parsed.blocks if "Автор:" in block.text)

    value = original[source.locator["char_start"]:source.locator["char_end"]]
    assert "Автор: Алексей &amp; команда" in value
    assert source.locator["line_start"] == source.locator["line_end"] == 3


def test_pdf_text_is_extracted_with_page_locator() -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
    })
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (Project owner is Team Alpha. The project uses Python.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    stream = BytesIO()
    writer.write(stream)

    parsed = parse_document("brief.pdf", stream.getvalue())

    assert parsed.metadata["page_count"] == 1
    assert parsed.metadata["page_width"] == 300
    assert parsed.metadata["page_height"] == 300
    assert parsed.blocks[0].locator["page"] == 1
    assert "Team Alpha" in parsed.blocks[0].text


def test_empty_pdf_page_is_not_mistaken_for_a_scan_and_encrypted_pdf_still_fails() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    stream = BytesIO()
    writer.write(stream)
    parsed = parse_document("scan.pdf", stream.getvalue())
    assert parsed.blocks == []
    assert parsed.metadata["ocr_required"] is False
    assert parsed.metadata["page_count"] == 1
    assert parsed.metadata["pdf_page_map"][0]["classification"] == "blank"

    writer.encrypt("password")
    encrypted = BytesIO()
    writer.write(encrypted)
    with pytest.raises(DocumentParsingError, match="паролем"):
        parse_document("locked.pdf", encrypted.getvalue())


def test_mixed_pdf_classifies_each_page_and_records_crop_rotation() -> None:
    fixture = Path(__file__).parent / "fixtures" / "mixed.pdf"

    parsed = parse_document("mixed.pdf", fixture.read_bytes())

    page_map = parsed.metadata["pdf_page_map"]
    assert [page["classification"] for page in page_map] == ["native", "ocr_candidate", "blank", "ocr_candidate"]
    assert parsed.metadata["ocr_pages"] == [2, 4]
    assert parsed.metadata["ocr_required"] is True
    assert {block.locator["page"] for block in parsed.blocks} == {1}
    assert page_map[3]["rotation"] == 90
    assert len(page_map[3]["crop_box"]) == 4


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("document.exe", b"text"),
        ("report.pdf", b"not a pdf"),
        ("broken.docx", b"not a zip archive"),
        ("broken.xml", b"<root>"),
        ("empty.txt", b""),
    ],
)
def test_invalid_inputs_fail_with_domain_error(filename: str, content: bytes) -> None:
    with pytest.raises(DocumentParsingError):
        parse_document(filename, content)
