from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services.document_search import (
    OriginalSearchSourceCache,
    find_text_matches,
    search_markdown,
    search_source_blocks,
    search_table_bytes,
)
from app.services.parsing import DocumentParsingError


def source(text: str, locator: dict, ordinal: int = 0, *, derived: bool = False):
    return SimpleNamespace(
        id=f"source-{ordinal}",
        text=text,
        locator=locator,
        ordinal=ordinal,
        is_derived=derived,
    )


def test_original_search_cache_is_size_bounded_lru_and_purges_deleted_documents() -> None:
    cache = OriginalSearchSourceCache(max_chars=10, max_entries=2)
    first = (source("1234", {"kind": "txt"}, 1),)
    second = (source("5678", {"kind": "txt"}, 2),)
    third = (source("9012", {"kind": "txt"}, 3),)
    cache.put(("doc-a", 1, "a"), first)
    cache.put(("doc-b", 1, "b"), second)

    assert cache.get(("doc-a", 1, "a")) is first
    cache.put(("doc-c", 1, "c"), third)
    assert cache.get(("doc-b", 1, "b")) is None
    assert cache.get(("doc-a", 1, "a")) is first
    assert cache.get(("doc-c", 1, "c")) is third

    cache.remove_document("doc-a")
    assert cache.get(("doc-a", 1, "a")) is None


def test_normalized_search_keeps_exact_ranges_for_case_spacing_and_cyrillic() -> None:
    value = "Тема:\tПРОВЕРКА\nдокумента и проверка"

    matches = find_text_matches(value, "проверка документа")

    assert [(item.start, item.end) for item in matches] == [(6, 24)]
    assert value[matches[0].start:matches[0].end] == "ПРОВЕРКА\nдокумента"


def test_search_scans_large_text_and_returns_codepoint_ranges_after_astral_characters() -> None:
    prefix = ("строка 🚀 текст без совпадения\n" * 4_000)
    source_text = prefix + ("фрагмент SEARCHABLE здесь\n" * 100)

    result = search_source_blocks(
        [source(source_text, {"kind": "txt", "source_text": source_text, "char_start": 0, "line_start": 1})],
        "searchable", file_type="txt", document_id="doc-long", processing_version=2,
        offset=99, limit=1,
    )

    assert result["total"] == 100
    match = result["matches"][0]
    start = len(prefix) + 99 * len("фрагмент SEARCHABLE здесь\n") + len("фрагмент ")
    assert match["match_start"] == start
    assert source_text[match["match_start"]:match["match_end"]] == "SEARCHABLE"
    assert match["locator"]["char_start"] == start
    assert match["locator"]["line_start"] == 4_000 + 100


def test_normalized_match_returns_overlapping_occurrences_without_repeating_source() -> None:
    assert [(item.start, item.end) for item in find_text_matches("aaaaa", "aa")] == [
        (0, 2), (1, 3), (2, 4), (3, 5),
    ]


def test_search_counts_full_source_and_returns_requested_page_without_derived_rows() -> None:
    chunks = [
        source("Целевое слово", {"kind": "txt", "source_text": "Целевое слово", "char_start": 10, "line_start": 2}, 0),
        source("Целевое слово", {"kind": "txt", "source_text": "Целевое слово", "char_start": 40, "line_start": 7}, 1),
        source("Целевое слово", {"kind": "txt", "source_text": "Целевое слово"}, 2, derived=True),
    ]

    result = search_source_blocks(
        chunks, "ЦЕЛЕВОЕ", file_type="txt", document_id="doc-1", processing_version=4, offset=1, limit=1,
    )

    assert result["total"] == 2
    assert len(result["matches"]) == 1
    assert result["matches"][0]["locator"]["char_start"] == 40
    assert result["matches"][0]["locator"]["processing_version"] == 4
    assert result["matches"][0]["text"] == "Целевое"


def test_search_returns_only_ocr_boxes_for_the_matching_words() -> None:
    chunk = source("Привет мир", {
        "kind": "pdf", "page": 5, "char_start": 100, "ocr": True, "source_text": "Привет мир",
        "ocr_map": {"coordinate_scale": 1000, "word_boxes": [
            [10, 10, 100, 50, 100, 106, 1, 92],
            [110, 10, 200, 50, 107, 110, 1, 88],
        ]},
    })

    result = search_source_blocks([chunk], "МИР", file_type="pdf", document_id="doc-1", processing_version=3)

    assert result["total"] == 1
    locator = result["matches"][0]["locator"]
    assert locator["page"] == 5
    assert locator["ocr_map"]["word_boxes"] == [[110, 10, 200, 50, 107, 110, 1, 88]]


@pytest.mark.parametrize(
    ("file_type", "locator", "preserved"),
    [
        ("pdf", {"kind": "pdf", "page": 4, "char_start": 100}, {"page": 4, "char_start": 109, "char_end": 115}),
        ("docx", {"kind": "docx", "paragraph": 8, "table": 2, "row": 5}, {"paragraph": 8, "table": 2, "row": 5}),
        ("xml", {"kind": "xml", "path": "/root/item[2]", "char_start": 40}, {"path": "/root/item[2]", "char_start": 49, "char_end": 55}),
        ("json", {"kind": "json", "path": "$.items[1].name", "char_start": 12}, {"path": "$.items[1].name", "char_start": 21, "char_end": 27}),
        ("html", {"kind": "html", "path": "body > p:nth-of-type(2)", "char_start": 70}, {"path": "body > p:nth-of-type(2)", "char_start": 79, "char_end": 85}),
        ("pptx", {"kind": "pptx", "slide": 3, "shape": 7}, {"slide": 3, "shape": 7}),
        ("epub", {"kind": "epub", "chapter": 2, "path": "OEBPS/chapter2.xhtml", "element": "p"}, {"chapter": 2, "path": "OEBPS/chapter2.xhtml", "element": "p"}),
    ],
)
def test_search_keeps_native_locators_for_each_document_family(file_type: str, locator: dict, preserved: dict) -> None:
    text = "Контекст TARGET после текста"
    result = search_source_blocks(
        [source(text, {**locator, "source_text": text}, 4)],
        "target",
        file_type=file_type,
        document_id="doc-locator",
        processing_version=6,
    )

    assert result["total"] == 1
    match = result["matches"][0]
    assert {key: match["locator"][key] for key in preserved} == preserved
    assert match["locator"]["search_match"] is True
    assert match["locator"]["search_range"] == {
        "coordinate_space": "source-text",
        "start": 9,
        "end": 15,
    }


def test_markdown_search_uses_global_codepoint_offsets_and_maps_to_active_chunk() -> None:
    text = "x\n" * 80 + "Строка для ПОИСКА\n" + "z\n" * 80
    match_start = text.index("ПОИСКА")
    chunk = SimpleNamespace(
        ordinal=12,
        markdown_char_start=match_start - 7,
        markdown_char_end=match_start + 6,
        locator={"kind": "docx", "paragraph": 19, "source_text": "Строка для ПОИСКА"},
    )

    result = search_markdown(
        text, "поиска", chunks=[chunk], document_id="doc-2", processing_version=2,
        file_type="docx", offset=0, limit=10,
    )

    assert result["total"] == 1
    hit = result["matches"][0]
    assert hit["text"] == "ПОИСКА"
    assert hit["markdown_start"] == match_start
    assert hit["locator"]["markdown_char_start"] == match_start
    assert hit["locator"]["markdown_line_start"] == 81
    assert hit["locator"]["paragraph"] == 19


def test_csv_search_resolves_quoted_cp1251_cells_and_paginates_global_results() -> None:
    data = "Название;Описание\r\nДокумент;\"Срок, утверждён; архив\"\r\nДокумент;утверждён повторно\r\n".encode("cp1251")

    result = search_table_bytes(data, "csv", "УТВЕРЖДЁН", offset=1, limit=1)

    assert result["total"] == 2
    hit = result["matches"][0]
    assert hit["text"] == "утверждён"
    assert hit["locator"]["row_start"] == 3
    assert hit["locator"]["column"] == "Описание"
    assert hit["locator"]["column_index"] == 1


def test_csv_search_keeps_header_cells_navigable() -> None:
    result = search_table_bytes("Название,Комментарий\nКнига,готово\n".encode(), "csv", "КОММЕНТАРИЙ")

    assert result["total"] == 1
    assert result["matches"][0]["locator"]["row_start"] == 1
    assert result["matches"][0]["locator"]["column"] == "Комментарий"


def test_xlsx_search_covers_later_sheets_and_returns_cell_coordinates() -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Кратко"
    first.append(["Имя", "Статус"])
    first.append(["Альфа", "Обычный"])
    second = workbook.create_sheet("Архив")
    second.append(["Документ", "Описание"])
    second.append(["Файл", "Редкая фраза для поиска"])
    data = BytesIO()
    workbook.save(data)

    result = search_table_bytes(data.getvalue(), "xlsx", "РЕДКАЯ ФРАЗА")

    assert result["total"] == 1
    hit = result["matches"][0]
    assert hit["locator"]["sheet"] == "Архив"
    assert hit["locator"]["row_start"] == 2
    assert hit["locator"]["column"] == "Описание"


def test_xlsx_search_preserves_physical_row_when_blank_rows_precede_a_hit() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Название", "Статус"])
    sheet.append(["", ""])
    sheet.append(["", ""])
    sheet.append(["Файл", "Уникальный маркер"])
    data = BytesIO()
    workbook.save(data)

    result = search_table_bytes(data.getvalue(), "xlsx", "маркер")

    assert result["total"] == 1
    assert result["matches"][0]["locator"]["row_start"] == 4


def test_xls_search_covers_sheet_physical_rows_and_cells() -> None:
    import xlwt

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Архив")
    for row, values in enumerate([
        ["Имя", "Описание"],
        ["Старый", "обычный текст"],
        ["Документ", "Уникальный маркер"],
    ]):
        for column, value in enumerate(values):
            sheet.write(row, column, value)
    data = BytesIO()
    workbook.save(data)

    result = search_table_bytes(data.getvalue(), "xls", "УНИКАЛЬНЫЙ МАРКЕР")

    assert result["total"] == 1
    assert result["matches"][0]["text"] == "Уникальный маркер"
    assert result["matches"][0]["locator"]["sheet"] == "Архив"
    assert result["matches"][0]["locator"]["row_start"] == 3
    assert result["matches"][0]["locator"]["column"] == "Описание"


def test_search_query_is_literal_and_limits_are_validated() -> None:
    matches = find_text_matches("value .* literal", ".*")
    assert [(match.start, match.end) for match in matches] == [(6, 8)]

    with pytest.raises(DocumentParsingError):
        search_source_blocks([], "   ", file_type="txt", document_id="doc", processing_version=1)
    with pytest.raises(DocumentParsingError):
        search_source_blocks([], "x", file_type="txt", document_id="doc", processing_version=1, limit=101)
