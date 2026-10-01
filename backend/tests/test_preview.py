from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import Workbook

from app.services.parsing import DocumentParsingError
from app.services.preview import build_preview, read_csv_table, read_spreadsheet_table


def chunk(text: str, locator: dict, ordinal: int = 0, *, derived: bool = False) -> SimpleNamespace:
    return SimpleNamespace(
        id=uuid4(),
        text=text,
        locator=locator,
        ordinal=ordinal,
        is_derived=derived,
    )


@pytest.mark.parametrize(
    ("file_type", "layout", "kind", "metadata"),
    [
        ("pdf", "pdf", "page", {"page_count": 2}),
        ("docx", "paper", "paragraph", {"paragraph_count": 1}),
        ("txt", "paper", "text", {"line_count": 2}),
        ("md", "paper", "text", {"line_count": 2}),
        ("csv", "table", "row", {"columns": ["Name", "Value"]}),
        ("xlsx", "table", "row", {"columns": ["Name", "Value"]}),
        ("xls", "table", "row", {"columns": ["Name", "Value"]}),
        ("pptx", "slides", "paragraph", {"slide_count": 1}),
        ("html", "paper", "text", {"line_count": 2}),
        ("json", "tree", "text", {"line_count": 2}),
        ("epub", "paper", "text", {"chapter_count": 1}),
        ("xml", "tree", "node", {"root": "report"}),
    ],
)
def test_preview_keeps_source_anchor_for_every_supported_format(
    file_type: str,
    layout: str,
    kind: str,
    metadata: dict,
) -> None:
    source = chunk("Evidence", {"kind": file_type, "label": "Источник 1"})

    preview = build_preview(
        document_id="document-1",
        file_type=file_type,
        metadata=metadata,
        chunks=[source],
        original_url="/api/v1/documents/document-1/file",
    )

    assert preview["layout"] == layout
    assert preview["blocks"][0]["kind"] == kind
    assert preview["blocks"][0]["source_id"] == str(source.id)
    locator = preview["blocks"][0]["locator"]
    assert {key: locator[key] for key in source.locator} == source.locator
    assert locator["locator_version"] == 1
    assert locator["document_id"] == "document-1"
    assert locator["processing_version"] == 1
    assert locator["source_type"]
    assert locator["source_range"]["coordinate_space"]
    assert {key: value for key, value in locator["source_range"].items() if key != "coordinate_space"} == {}
    assert preview["aspect_ratio"] > 0
    assert preview["total_blocks"] == 1
    assert preview["truncated"] is False
    assert preview["original_url"] == "/api/v1/documents/document-1/file"
    assert preview["renderer"] == {
        "pdf": "pdf", "docx": "docx", "txt": "text", "md": "text", "csv": "csv", "xml": "xml",
        "xlsx": "xlsx", "xls": "xls", "pptx": "pptx", "html": "html", "json": "json", "epub": "epub",
    }[file_type]
    assert preview["source_count"] == 1


def test_csv_preview_recovers_cells_without_delegating_table_shape_to_model() -> None:
    source = chunk(
        "Name: Alpha | Value: 10\nName: Beta | Value: 20",
        {"kind": "csv", "label": "Строки 2–3", "row_start": 2, "row_end": 3},
    )

    preview = build_preview(
        document_id="document-1",
        file_type="csv",
        metadata={"columns": ["Name", "Value"]},
        chunks=[source],
        original_url="/api/v1/documents/document-1/file",
    )

    assert preview["blocks"][0]["rows"] == [["Alpha", "10"], ["Beta", "20"]]


def test_preview_uses_source_page_dimensions_for_pdf_and_docx() -> None:
    source = chunk("Evidence", {"kind": "pdf", "page": 1})

    pdf = build_preview(
        document_id="document-1",
        file_type="pdf",
        metadata={"page_count": 1, "page_width": 100, "page_height": 200},
        chunks=[source],
        original_url="/api/v1/documents/document-1/file",
    )
    docx = build_preview(
        document_id="document-1",
        file_type="docx",
        metadata={"page_width": 297, "page_height": 210},
        chunks=[source],
        original_url=None,
    )

    assert pdf["aspect_ratio"] == 0.5
    assert docx["aspect_ratio"] == 297 / 210


def test_derived_csv_calculation_is_marked_and_has_no_fake_cells() -> None:
    source = chunk(
        "Локальный расчёт: сумма 30",
        {"kind": "csv_derived", "label": "Показатели столбца «Value»", "derived": True},
        derived=True,
    )

    preview = build_preview(
        document_id="document-1",
        file_type="csv",
        metadata={"columns": ["Name", "Value"]},
        chunks=[source],
        original_url=None,
    )

    assert preview["blocks"][0]["kind"] == "calculation"
    assert preview["blocks"][0]["rows"] is None


def test_preview_marks_large_documents_as_truncated() -> None:
    sources = [chunk(f"Block {index}", {"kind": "txt"}, index) for index in range(2_005)]

    preview = build_preview(
        document_id="document-1",
        file_type="txt",
        metadata={},
        chunks=sources,
        original_url=None,
    )

    assert len(preview["blocks"]) == 2_000
    assert preview["total_blocks"] == 2_005
    assert preview["truncated"] is True


def test_preview_uses_index_count_when_api_fetches_only_a_safe_prefix() -> None:
    source = chunk("Block", {"kind": "txt"})

    preview = build_preview(
        document_id="document-1",
        file_type="txt",
        metadata={},
        chunks=[source],
        original_url=None,
        total_blocks=4_500,
    )

    assert preview["total_blocks"] == 4_500
    assert preview["truncated"] is True


@pytest.mark.parametrize(
    ("delimiter", "payload"),
    [
        (",", "Name,Value\nAlpha,10\nBeta,20\n"),
        (";", "Name;Value\nAlpha;10\nBeta;20\n"),
        ("\t", "Name\tValue\nAlpha\t10\nBeta\t20\n"),
    ],
)
def test_original_csv_table_preserves_delimiters_and_source_row_numbers(delimiter: str, payload: str) -> None:
    table = read_csv_table(payload.encode("utf-8"), offset=1, limit=1)

    assert table["columns"] == ["Name", "Value"]
    assert table["rows"] == [{"number": 3, "cells": ["Beta", "20"]}]
    assert table["total_rows"] == 2
    assert table["delimiter"] == delimiter


def test_original_csv_table_decodes_cp1251_and_caps_page_size() -> None:
    table = read_csv_table("Название;Значение\nТест;готово\n".encode("cp1251"), limit=9999)

    assert table["columns"] == ["Название", "Значение"]
    assert table["rows"][0]["cells"] == ["Тест", "готово"]
    assert table["limit"] == 500


def test_original_xlsx_table_selects_sheet_and_keeps_physical_row_numbers() -> None:
    workbook = Workbook()
    workbook.active.title = "Обзор"
    workbook.active.append(["Title", "Value"])
    workbook.active.append(["Overview row", 1])
    worksheet = workbook.create_sheet("Источники")
    worksheet.append(["Title", "Value"])
    worksheet.append(["First", 10])
    worksheet.append(["Second", 20])
    worksheet.append(["Third", 30])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    table = read_spreadsheet_table(stream.getvalue(), "xlsx", sheet="Источники", offset=1, limit=1)

    assert table["sheet"] == "Источники"
    assert table["available_sheets"] == ["Обзор", "Источники"]
    assert table["rows"] == [{"number": 3, "cells": ["Second", "20"]}]
    assert table["total_rows"] == 3


def test_original_xlsx_table_rejects_unknown_sheet() -> None:
    workbook = Workbook()
    workbook.active.append(["Value"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    with pytest.raises(DocumentParsingError, match="лист книги не найден"):
        read_spreadsheet_table(stream.getvalue(), "xlsx", sheet="Missing")
