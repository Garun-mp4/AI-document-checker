"""Generate synthetic fixtures; no user documents are read."""
import importlib.util
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "frontend/e2e/fixtures"
spec = importlib.util.spec_from_file_location("base_fixtures", ROOT / "backend/tests/fixtures/generate_fixtures.py")
base = importlib.util.module_from_spec(spec)
spec.loader.exec_module(base)


def generate():
    OUT.mkdir(parents=True, exist_ok=True)
    contents = dict(base.FIXTURE_CONTENTS)
    contents["sample.htm"] = contents["sample.html"]
    for ext in ("txt", "md"):
        contents[f"cp1251.{ext}"] = contents[f"sample.{ext}"].decode().encode("cp1251")
    contents["cp1251.csv"] = 'Проект;Часы\n"Альфа; Бета";10\nГамма;20\n'.encode("cp1251")
    contents["quoted.csv"] = 'Проект,Описание,Часы\nАльфа,"запятая, кавычки ""да""",10\nБета,"две\nстроки",20\n'.encode()
    contents["tab.csv"] = 'Проект\tЧасы\nАльфа\t10\nБета\t20\n'.encode()
    contents["long.txt"] = ("Цель: проверить длинные строки.\n" + "Длинная строка " * 180 + "\nАвтор: Алексей Пример.\n").encode()
    contents["large.csv"] = ("Проект,Часы\n" + "".join(f"Проект {i},{i}\n" for i in range(1, 241))).encode()
    pdf = PdfReader(BytesIO(contents["sample.pdf"]))
    writer = PdfWriter()
    for _ in range(3):
        writer.add_page(pdf.pages[0])
    stream = BytesIO()
    writer.write(stream)
    contents["multipage.pdf"] = stream.getvalue()
    writer.encrypt("synthetic-secret")
    stream = BytesIO()
    writer.write(stream)
    contents["encrypted.pdf"] = stream.getvalue()
    document = Document(BytesIO(contents["sample.docx"]))
    document.add_page_break()
    document.add_paragraph("Вторая страница: дополнительный источник, срок 2026-12-01.")
    stream = BytesIO()
    document.save(stream)
    contents["multipage.docx"] = stream.getvalue()
    workbook = Workbook()
    for sheet in (workbook.active, workbook.create_sheet("Второй лист")):
        sheet.append(["Проект", "Часы"])
        sheet.append(["Альфа", 10])
        sheet.append(["Бета", 20])
    stream = BytesIO()
    workbook.save(stream)
    contents["multisheet.xlsx"] = stream.getvalue()
    import xlwt
    book = xlwt.Workbook()
    for title in ("Первый", "Второй"):
        sheet = book.add_sheet(title)
        for row, values in enumerate([["Проект", "Часы"], ["Альфа", 10], ["Бета", 20]]):
            for column, value in enumerate(values):
                sheet.write(row, column, value)
    stream = BytesIO()
    book.save(stream)
    contents["multisheet.xls"] = stream.getvalue()
    presentation = Presentation(BytesIO(contents["sample.pptx"]))
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = "Второй слайд: срок 2026-12-01."
    stream = BytesIO()
    presentation.save(stream)
    contents["multislide.pptx"] = stream.getvalue()
    stream = BytesIO()
    with ZipFile(BytesIO(contents["sample.epub"])) as original, ZipFile(stream, "w") as target:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename.endswith("content.opf"):
                data = data.decode().replace("</manifest>", '<item id="second" href="chapter2.xhtml" media-type="application/xhtml+xml"/></manifest>').replace("</spine>", '<itemref idref="second"/></spine>').encode()
            target.writestr(item, data)
        target.writestr("OEBPS/chapter2.xhtml", "<html><body><h1>Вторая глава</h1><p>Срок: 2026-12-01.</p></body></html>")
    contents["multichapter.epub"] = stream.getvalue()
    contents["unsafe.html"] = '<html><body><h1>Тестовый проект</h1><script>window.__unsafeExecuted=true</script><img src="https://example.invalid/tracker" onerror="window.__unsafeExecuted=true"><p>Автор: Алексей Пример.</p></body></html>'.encode()
    contents["dangerous.xml"] = b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><x>&secret;</x>'
    # Synthetic hostile archives remain small on disk; none contains user data.
    for extension in ('docx', 'xlsx', 'pptx', 'epub'):
        stream = BytesIO()
        with ZipFile(BytesIO(contents[f'sample.{extension}'])) as original, ZipFile(stream, 'w', ZIP_DEFLATED) as target:
            for item in original.infolist():
                target.writestr(item, original.read(item.filename))
            target.writestr('../outside.txt', 'must not be extracted')
        contents[f'traversal.{extension}'] = stream.getvalue()
    for name, target_url in [('remote.docx', 'https://example.invalid/image.png'), ('unsafe-link.docx', 'file:///etc/passwd')]:
        stream = BytesIO()
        with ZipFile(BytesIO(contents['sample.docx'])) as original, ZipFile(stream, 'w', ZIP_DEFLATED) as target:
            for item in original.infolist():
                if item.filename != 'word/_rels/document.xml.rels':
                    target.writestr(item, original.read(item.filename))
            target.writestr('word/_rels/document.xml.rels', f'<Relationships><Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="{target_url}" TargetMode="External"/></Relationships>')
        contents[name] = stream.getvalue()
    for name, extension, extra in [
        ('zip-bomb.docx', 'docx', [('big.bin', b'x' * (32 * 1024 * 1024 + 1))]),
        ('many-parts.epub', 'epub', [(f'part{i}.bin', b'') for i in range(5001)]),
        ('entity.xlsx', 'xlsx', [('evil.xml', b'<!DOCTYPE x [<!ENTITY s SYSTEM "file:///etc/passwd">]><x>&s;</x>')]),
    ]:
        stream = BytesIO()
        with ZipFile(BytesIO(contents[f'sample.{extension}'])) as original, ZipFile(stream, 'w', ZIP_DEFLATED) as target:
            for item in original.infolist():
                target.writestr(item, original.read(item.filename))
            for member, value in extra:
                target.writestr(member, value)
        contents[name] = stream.getvalue()
    contents['binary.txt'] = b'MZ disguised executable'
    contents["empty.txt"] = b""
    contents["corrupt.pdf"] = b"%PDF-1.7\ncorrupt"
    contents["corrupt.docx"] = b"PK\x03\x04corrupt"
    contents["wrong.pdf"] = b"This is not a PDF"
    contents["unsupported.exe"] = b"synthetic"
    contents["oversized.txt"] = b"x" * (25 * 1024 * 1024 + 1)
    blank = Image.new("RGB", (1240, 1754), "white")
    stream = BytesIO()
    blank.save(stream, "PDF", resolution=150)
    contents["unreadable.pdf"] = stream.getvalue()
    draw = ImageDraw.Draw(blank)
    fonts = [Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    font_path = next(path for path in fonts if path.exists())
    font = ImageFont.truetype(str(font_path), 38)
    for index, line in enumerate(["Тестовый проект OCR", "Автор: Алексей Пример", "Цель: проверить распознавание скана", "Срок: 30 ноября 2026 года"]):
        draw.text((70, 140 + index * 80), line, font=font, fill="black")
    stream = BytesIO()
    blank.save(stream, "PDF", resolution=150)
    contents["scan.pdf"] = stream.getvalue()
    for name, data in contents.items():
        (OUT / name).write_bytes(data)
    (OUT / "manifest.json").write_text(json.dumps({name: len(data) for name, data in contents.items()}, indent=2), encoding="utf-8")
    print(f"Generated {len(contents)} synthetic files in {OUT}")


if __name__ == "__main__":
    generate()
