"""Server-side querying and auditable calculations for uploaded tables."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from decimal import ROUND_HALF_UP, Decimal, localcontext
from typing import Any

from app.services.numeric_values import parse_decimal

MAX_TABLE_ROWS = 500_000
MAX_TABLE_COLUMNS = 500
MAX_TABLE_NONEMPTY_CELLS = 2_000_000


@dataclass(frozen=True)
class FormulaCell:
    column_index: int
    formula: str
    has_cached_value: bool


@dataclass(frozen=True)
class TableRow:
    number: int
    cells: tuple[str, ...]
    formulas: tuple[FormulaCell, ...] = ()

    def value(self, column_index: int) -> str:
        return self.cells[column_index] if column_index < len(self.cells) else ""

    def has_formula(self, column_index: int) -> bool:
        return any(formula.column_index == column_index for formula in self.formulas)


@dataclass(frozen=True)
class TableDataset:
    columns: tuple[str, ...]
    rows: tuple[TableRow, ...]
    sheet: str | None
    available_sheets: tuple[str, ...]
    formula_policy: str
    delimiter: str | None = None


def _error(message: str) -> None:
    from app.services.parsing import DocumentParsingError

    raise DocumentParsingError(message)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds") if value.time() != time.min else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _trim_cells(cells: list[str], formulas: list[FormulaCell] | None = None) -> tuple[str, ...]:
    last_value = max((index for index, cell in enumerate(cells) if cell.strip()), default=-1)
    last_formula = max((item.column_index for item in formulas or ()), default=-1)
    return tuple(cells[:max(last_value, last_formula) + 1])


def _validate_size(row_count: int, column_count: int, nonempty_cells: int) -> None:
    if row_count > MAX_TABLE_ROWS or column_count > MAX_TABLE_COLUMNS:
        _error("Таблица превышает безопасный предел: 500 000 строк или 500 столбцов.")
    if nonempty_cells > MAX_TABLE_NONEMPTY_CELLS:
        _error("В таблице слишком много заполненных ячеек для безопасной сортировки и расчётов.")


def detect_csv_dialect(text: str) -> tuple[csv.Dialect, str]:
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        return dialect, dialect.delimiter
    except csv.Error:
        delimiter = ";" if text.count(";") > text.count(",") else ","
        return csv.excel, delimiter


def _read_csv(data: bytes) -> TableDataset:
    from app.services.parsing import DocumentParsingError, _decode_text_with_encoding

    text, _encoding = _decode_text_with_encoding(data)
    dialect, delimiter = detect_csv_dialect(text)
    try:
        records = list(csv.reader(io.StringIO(text, newline=""), dialect=dialect, delimiter=delimiter))
    except csv.Error as exc:
        raise DocumentParsingError("Не удалось разобрать строки CSV.") from exc
    header_index = next((index for index, row in enumerate(records) if any(cell.strip() for cell in row)), None)
    if header_index is None:
        raise DocumentParsingError("CSV пустой — строк не найдено.")
    raw_headers = records[header_index]
    if len(raw_headers) > MAX_TABLE_COLUMNS:
        raise DocumentParsingError("В CSV слишком много столбцов (максимум 500).")
    columns = tuple(cell.strip() or f"Столбец {index + 1}" for index, cell in enumerate(raw_headers))
    data_records = records[header_index + 1:]
    if len(data_records) > MAX_TABLE_ROWS:
        raise DocumentParsingError("CSV превышает безопасный предел в 500 000 строк.")
    rows: list[TableRow] = []
    nonempty_cells = sum(bool(value.strip()) for value in raw_headers)
    for record_index, record in enumerate(data_records, start=header_index + 2):
        cells = [str(value) for value in record[:len(columns)]]
        nonempty_cells += sum(bool(value.strip()) for value in cells)
        _validate_size(len(data_records), len(columns), nonempty_cells)
        rows.append(TableRow(record_index, _trim_cells(cells)))
    return TableDataset(columns, tuple(rows), None, (), "not_applicable", delimiter)


def _read_xlsx(data: bytes, sheet: str | None) -> TableDataset:
    from app.services.parsing import DocumentParsingError

    try:
        from openpyxl import load_workbook
        formulas_book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        values_book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except ImportError as exc:
        raise DocumentParsingError("Для XLSX не установлен модуль openpyxl.") from exc
    except Exception as exc:
        raise DocumentParsingError("XLSX повреждён или имеет неверную структуру.") from exc
    try:
        available_sheets = list(values_book.sheetnames)
        selected_name = sheet or (available_sheets[0] if available_sheets else None)
        if selected_name not in available_sheets:
            raise DocumentParsingError("Выбранный лист книги не найден.")
        formula_sheet = formulas_book[selected_name]
        value_sheet = values_book[selected_name]
        row_count = formula_sheet.max_row or 0
        column_count = formula_sheet.max_column or 0
        if row_count > MAX_TABLE_ROWS + 1 or column_count > MAX_TABLE_COLUMNS:
            raise DocumentParsingError("Размер листа превышает безопасный предел строк или столбцов.")
        if row_count == 0 or column_count == 0:
            raise DocumentParsingError("Таблица пустая — столбцы не найдены.")
        formula_rows = formula_sheet.iter_rows(min_row=1, max_row=row_count, max_col=column_count, values_only=False)
        value_rows = value_sheet.iter_rows(min_row=1, max_row=row_count, max_col=column_count, values_only=False)
        header_values = next(value_rows)
        columns = tuple(_text(cell.value).strip() or f"Столбец {index + 1}" for index, cell in enumerate(header_values))
        next(formula_rows)
        rows: list[TableRow] = []
        nonempty_cells = sum(bool(value.strip()) for value in columns)
        for physical_row, (formula_row, value_row) in enumerate(zip(formula_rows, value_rows, strict=True), start=2):
            formula_cells: list[FormulaCell] = []
            cells: list[str] = []
            for column_index, (formula_cell, value_cell) in enumerate(zip(formula_row, value_row, strict=True)):
                formula = formula_cell.value
                is_formula = formula_cell.data_type == "f"
                cached_value = value_cell.value
                if is_formula:
                    formula_text = str(getattr(formula, "text", formula))
                    formula_cells.append(FormulaCell(
                        column_index,
                        formula_text if formula_text.startswith("=") else f"={formula_text}",
                        cached_value is not None,
                    ))
                cells.append(_text(cached_value))
            nonempty_cells += sum(bool(value.strip()) for value in cells) + len(formula_cells)
            _validate_size(row_count - 1, column_count, nonempty_cells)
            rows.append(TableRow(physical_row, _trim_cells(cells, formula_cells), tuple(formula_cells)))
        return TableDataset(columns, tuple(rows), selected_name, tuple(available_sheets), "detected")
    except DocumentParsingError:
        raise
    except Exception as exc:
        raise DocumentParsingError("XLSX повреждён или имеет неверную структуру.") from exc
    finally:
        formulas_book.close()
        values_book.close()


def _read_xls(data: bytes, sheet: str | None) -> TableDataset:
    from app.services.parsing import DocumentParsingError

    try:
        import xlrd
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    except ImportError as exc:
        raise DocumentParsingError("Для XLS не установлен модуль xlrd.") from exc
    except Exception as exc:
        raise DocumentParsingError("XLS повреждён или имеет неверную структуру.") from exc
    try:
        available_sheets = workbook.sheet_names()
        selected_name = sheet or (available_sheets[0] if available_sheets else None)
        if selected_name not in available_sheets:
            raise DocumentParsingError("Выбранный лист книги не найден.")
        worksheet = workbook.sheet_by_name(selected_name)
        if worksheet.nrows > MAX_TABLE_ROWS + 1 or worksheet.ncols > MAX_TABLE_COLUMNS:
            raise DocumentParsingError("Размер листа превышает безопасный предел строк или столбцов.")
        if worksheet.nrows == 0 or worksheet.ncols == 0:
            raise DocumentParsingError("Таблица пустая — столбцы не найдены.")
        columns = tuple(_text(value).strip() or f"Столбец {index + 1}" for index, value in enumerate(worksheet.row_values(0)))
        rows: list[TableRow] = []
        nonempty_cells = sum(bool(value.strip()) for value in columns)
        for row_index in range(1, worksheet.nrows):
            cells: list[str] = []
            for column_index in range(worksheet.ncols):
                value = worksheet.cell_value(row_index, column_index)
                if worksheet.cell_type(row_index, column_index) == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                cells.append(_text(value))
            nonempty_cells += sum(bool(value.strip()) for value in cells)
            _validate_size(worksheet.nrows - 1, worksheet.ncols, nonempty_cells)
            rows.append(TableRow(row_index + 1, _trim_cells(cells)))
        return TableDataset(columns, tuple(rows), selected_name, tuple(available_sheets), "cached_values_only")
    except DocumentParsingError:
        raise
    except Exception as exc:
        raise DocumentParsingError("XLS повреждён или имеет неверную структуру.") from exc
    finally:
        if hasattr(workbook, "release_resources"):
            workbook.release_resources()


def read_table_dataset(data: bytes, file_type: str, *, sheet: str | None = None) -> TableDataset:
    if file_type == "csv":
        return _read_csv(data)
    if file_type == "xlsx":
        return _read_xlsx(data, sheet)
    if file_type == "xls":
        return _read_xls(data, sheet)
    _error("Табличный просмотр доступен только для CSV, XLSX и XLS.")


def _date_key(value: str) -> datetime | None:
    candidate = value.strip()
    if not candidate:
        return None
    parsed: datetime | None = None
    try:
        parsed = datetime.fromisoformat(candidate.replace("Z", "+00:00"))
    except ValueError:
        for pattern in ("%d.%m.%Y", "%d.%m.%Y %H:%M:%S", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                parsed = datetime.strptime(candidate, pattern).replace(tzinfo=timezone.utc)
                break
            except ValueError:
                continue
    if parsed is None:
        return None
    if parsed.tzinfo is not None:
        parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed


def _column_kind(rows: tuple[TableRow, ...], column_index: int) -> str:
    values = [row.value(column_index).strip() for row in rows]
    values = [value for value in values if value]
    if not values:
        return "empty"
    if all(_date_key(value) is not None for value in values):
        return "date"
    if all(parse_decimal(value) is not None for value in values):
        return "number"
    return "text"


def _filter_rows(rows: tuple[TableRow, ...], columns: tuple[str, ...], filter_spec: dict[str, Any] | None) -> list[TableRow]:
    if filter_spec is None:
        return list(rows)
    column_index = filter_spec["column_index"]
    if column_index < 0 or column_index >= len(columns):
        _error("Выбранный столбец фильтра не найден.")
    kind = filter_spec["kind"]
    operator = filter_spec["operator"]
    value = str(filter_spec.get("value") or "")
    if kind == "text" and operator not in {"contains", "equals"}:
        _error("Для текстового фильтра выберите «содержит» или «совпадает».")
    if kind == "number" and operator not in {"equals", "gt", "gte", "lt", "lte"}:
        _error("Для числового фильтра выберите знак сравнения.")
    if kind == "empty" and operator not in {"is_empty", "is_not_empty"}:
        _error("Для фильтра пустых значений выберите допустимое условие.")
    if kind not in {"text", "number", "empty"}:
        _error("Тип фильтра не поддерживается.")
    expected_number = parse_decimal(value) if kind == "number" else None
    if kind == "number" and expected_number is None:
        _error("Введите число: запятая или точка задаёт десятичную часть, пробелы — разряды тысяч.")
    if kind == "text" and not value.strip():
        _error("Введите текст для фильтра.")

    def matches(row: TableRow) -> bool:
        cell = row.value(column_index)
        nonempty = bool(cell.strip()) or row.has_formula(column_index)
        if kind == "empty":
            return not nonempty if operator == "is_empty" else nonempty
        if kind == "text":
            actual = cell.casefold()
            expected = value.casefold()
            return expected in actual if operator == "contains" else actual == expected
        actual_number = parse_decimal(cell)
        if actual_number is None or expected_number is None:
            return False
        return {
            "equals": actual_number == expected_number,
            "gt": actual_number > expected_number,
            "gte": actual_number >= expected_number,
            "lt": actual_number < expected_number,
            "lte": actual_number <= expected_number,
        }[operator]

    return [row for row in rows if matches(row)]


def _normalize_filter(
    filter_column: int | None,
    filter_kind: str | None,
    filter_operator: str | None,
    filter_value: str | None,
) -> dict[str, Any] | None:
    supplied = (filter_column is not None, filter_kind is not None, filter_operator is not None)
    if not any(supplied):
        if filter_value not in {None, ""}:
            _error("Укажите столбец, тип и условие фильтра.")
        return None
    if not all(supplied):
        _error("Для фильтра нужны столбец, тип и условие.")
    value = filter_value or ""
    if len(value) > 256:
        _error("Значение фильтра не должно превышать 256 символов.")
    return {"column_index": filter_column, "kind": filter_kind, "operator": filter_operator, "value": value}


def query_table(
    data: bytes,
    file_type: str,
    *,
    offset: int = 0,
    limit: int = 100,
    sheet: str | None = None,
    sort_column: int | None = None,
    sort_direction: str = "asc",
    filter_column: int | None = None,
    filter_kind: str | None = None,
    filter_operator: str | None = None,
    filter_value: str | None = None,
    focus_row: int | None = None,
) -> dict[str, Any]:
    if offset < 0:
        _error("offset не может быть отрицательным.")
    dataset = read_table_dataset(data, file_type, sheet=sheet)
    if sort_direction not in {"asc", "desc"}:
        _error("Направление сортировки не поддерживается.")
    if sort_column is not None and sort_column >= len(dataset.columns):
        _error("Столбец сортировки не найден.")
    limit = max(1, min(limit, 500))
    filter_spec = _normalize_filter(filter_column, filter_kind, filter_operator, filter_value)
    filtered_rows = _filter_rows(dataset.rows, dataset.columns, filter_spec)
    if sort_column is not None:
        nonempty = [row for row in filtered_rows if row.value(sort_column).strip()]
        uncached_formulas = [row for row in filtered_rows if not row.value(sort_column).strip() and row.has_formula(sort_column)]
        empty = [row for row in filtered_rows if not row.value(sort_column).strip() and not row.has_formula(sort_column)]
        mode = _column_kind(tuple(nonempty), sort_column)

        def sort_value(row: TableRow) -> Any:
            value = row.value(sort_column).strip()
            if mode == "number":
                return parse_decimal(value)
            if mode == "date":
                return _date_key(value)
            return value.casefold()

        nonempty.sort(key=sort_value, reverse=sort_direction == "desc")
        filtered_rows = nonempty + uncached_formulas + empty
    focus_visible: bool | None = None
    if focus_row is not None:
        focus_index = next((index for index, row in enumerate(filtered_rows) if row.number == focus_row), None)
        focus_visible = focus_index is not None
        if focus_index is not None:
            offset = max(0, focus_index - min(20, limit // 2))
    page_rows = filtered_rows[offset:offset + limit]
    return {
        "columns": list(dataset.columns),
        "rows": [{
            "number": row.number,
            "cells": [row.value(index) for index in range(len(dataset.columns))],
            "formula_cells": [{
                "column_index": formula.column_index,
                "formula": formula.formula,
                "has_cached_value": formula.has_cached_value,
            } for formula in row.formulas],
        } for row in page_rows],
        "offset": offset,
        "limit": limit,
        "total_rows": len(dataset.rows),
        "filtered_rows": len(filtered_rows),
        "sheet": dataset.sheet,
        "available_sheets": list(dataset.available_sheets),
        "delimiter": dataset.delimiter,
        "column_kinds": [_column_kind(dataset.rows, index) for index in range(len(dataset.columns))],
        "formula_policy": dataset.formula_policy,
        "sort_column": sort_column,
        "sort_direction": sort_direction if sort_column is not None else None,
        "filter": filter_spec,
        "focus_row": focus_row,
        "focus_row_visible": focus_visible,
    }


def _format_decimal(value: Decimal) -> str:
    return format(value, "f")


def _aggregate(rows: list[TableRow], column_index: int) -> dict[str, Any]:
    nonempty_rows = [row for row in rows if row.value(column_index).strip() or row.has_formula(column_index)]
    parsed = [(row, parse_decimal(row.value(column_index))) for row in nonempty_rows]
    numbers = [number for _row, number in parsed if number is not None]
    missing_formula_cache_count = sum(
        1 for row in nonempty_rows for formula in row.formulas
        if formula.column_index == column_index and not formula.has_cached_value
    )
    nonnumeric_count = sum(1 for _row, number in parsed if number is None) - missing_formula_cache_count
    result: dict[str, Any] = {
        "count": len(rows),
        "non_empty_count": len(nonempty_rows),
        "numeric_count": len(numbers),
        "nonnumeric_count": max(0, nonnumeric_count),
        "formula_count": sum(1 for row in rows for formula in row.formulas if formula.column_index == column_index),
        "formula_cache_missing_count": missing_formula_cache_count,
        "sum": None,
        "average": None,
        "minimum": None,
        "maximum": None,
        "source_row_start": min((row.number for row in rows), default=None),
        "source_row_end": max((row.number for row in rows), default=None),
    }
    if not numbers:
        return result
    scale = max((max(0, -number.as_tuple().exponent) for number in numbers), default=0)
    integer_digits = max((max(1, number.adjusted() + 1) for number in numbers), default=1)
    precision = max(28, integer_digits + scale + len(str(len(numbers))) + 4)
    with localcontext() as context:
        context.prec = precision
        total = sum(numbers, Decimal(0))
        average = (total / Decimal(len(numbers))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    result.update({
        "sum": _format_decimal(total),
        "average": _format_decimal(average),
        "minimum": _format_decimal(min(numbers)),
        "maximum": _format_decimal(max(numbers)),
    })
    return result


def calculate_table(
    data: bytes,
    file_type: str,
    *,
    sheet: str | None,
    column_index: int,
    filter_spec: dict[str, Any] | None = None,
) -> dict[str, Any]:
    dataset = read_table_dataset(data, file_type, sheet=sheet)
    if column_index < 0 or column_index >= len(dataset.columns):
        _error("Выбранный столбец расчёта не найден.")
    normalized_filter = None
    if filter_spec is not None:
        normalized_filter = _normalize_filter(
            filter_spec.get("column_index"), filter_spec.get("kind"),
            filter_spec.get("operator"), filter_spec.get("value"),
        )
    filtered_rows = _filter_rows(dataset.rows, dataset.columns, normalized_filter)
    return {
        "sheet": dataset.sheet,
        "available_sheets": list(dataset.available_sheets),
        "column_index": column_index,
        "column": dataset.columns[column_index],
        "filter": normalized_filter,
        "formula_policy": dataset.formula_policy,
        "rounding_rule": "Среднее округляется до 2 знаков по правилу ROUND_HALF_UP; сумма, минимум и максимум сохраняются без округления.",
        "document": {**_aggregate(list(dataset.rows), column_index), "scope": "document", "source_row_count": len(dataset.rows)},
        "filtered": {**_aggregate(filtered_rows, column_index), "scope": "current_filter", "source_row_count": len(filtered_rows)},
    }
