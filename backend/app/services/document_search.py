"""Bounded, locator-aware text search for original documents and Markdown."""
from __future__ import annotations

import csv
import io
import unicodedata
from bisect import bisect_right
from collections import OrderedDict
from collections.abc import Iterable
from dataclasses import dataclass
from threading import RLock
from typing import Any

from app.services.parsing import DocumentParsingError, _decode_text_with_encoding
from app.services.source_locators import versioned_source_locator

MAX_SEARCH_QUERY_CHARS = 256
MAX_SEARCH_PAGE_SIZE = 100
MAX_CACHED_SOURCE_CHARS = 8_000_000
MAX_CACHED_SOURCE_DOCUMENTS = 3


class OriginalSearchSourceCache:
    """Small process-local LRU for parsed search sources, bounded by text size."""

    def __init__(self, *, max_chars: int = MAX_CACHED_SOURCE_CHARS, max_entries: int = MAX_CACHED_SOURCE_DOCUMENTS):
        self.max_chars = max_chars
        self.max_entries = max_entries
        self._entries: OrderedDict[tuple[str, int, str], tuple[tuple[Any, ...], int]] = OrderedDict()
        self._chars = 0
        self._lock = RLock()

    def get(self, key: tuple[str, int, str]) -> tuple[Any, ...] | None:
        with self._lock:
            cached = self._entries.get(key)
            if cached is None:
                return None
            self._entries.move_to_end(key)
            return cached[0]

    def put(self, key: tuple[str, int, str], sources: Iterable[Any]) -> None:
        snapshot = tuple(sources)
        weight = sum(len(getattr(source, "text", "")) for source in snapshot)
        with self._lock:
            previous = self._entries.pop(key, None)
            if previous is not None:
                self._chars -= previous[1]
            if weight > self.max_chars or self.max_entries < 1:
                return
            while self._entries and (len(self._entries) >= self.max_entries or self._chars + weight > self.max_chars):
                _, (_, removed_weight) = self._entries.popitem(last=False)
                self._chars -= removed_weight
            self._entries[key] = (snapshot, weight)
            self._chars += weight

    def clear(self) -> None:
        with self._lock:
            self._entries.clear()
            self._chars = 0

    def remove_document(self, document_id: str) -> None:
        with self._lock:
            for key in [key for key in self._entries if key[0] == document_id]:
                _, weight = self._entries.pop(key)
                self._chars -= weight


original_search_source_cache = OriginalSearchSourceCache()


@dataclass(frozen=True)
class TextMatch:
    start: int
    end: int


def _normalise(value: str) -> str:
    result: list[str] = []
    whitespace = False
    for character in value:
        transformed = unicodedata.normalize("NFKC", character).casefold()
        if character.isspace() or transformed.isspace():
            if result and not whitespace:
                result.append(" ")
            whitespace = True
            continue
        whitespace = False
        result.extend(transformed)
    return "".join(result).strip()


def find_text_matches(value: str, query: str) -> list[TextMatch]:
    """Match normalized text while retaining offsets into the unmodified input."""
    return list(_iter_text_matches(value, query))


def _iter_text_matches(value: str, query: str):
    needle = _normalise(query)
    if not needle:
        return

    prefix = [0] * len(needle)
    candidate = 0
    for index in range(1, len(needle)):
        while candidate and needle[index] != needle[candidate]:
            candidate = prefix[candidate - 1]
        if needle[index] == needle[candidate]:
            candidate += 1
            prefix[index] = candidate

    recent: list[tuple[int, int]] = [(0, 0)] * len(needle)
    cursor = 0
    matched = 0

    def emit(character: str, start: int, end: int):
        nonlocal matched, cursor
        while matched and character != needle[matched]:
            matched = prefix[matched - 1]
        if character == needle[matched]:
            matched += 1
        recent[cursor] = (start, end)
        cursor = (cursor + 1) % len(needle)
        if matched == len(needle):
            yield TextMatch(recent[cursor][0], recent[(cursor - 1) % len(needle)][1])
            matched = prefix[matched - 1]

    pending_space: tuple[int, int] | None = None
    has_output = False
    for offset, character in enumerate(value):
        transformed = unicodedata.normalize("NFKC", character).casefold()
        if character.isspace() or transformed.isspace():
            if has_output:
                pending_space = (pending_space[0], offset + 1) if pending_space else (offset, offset + 1)
            continue
        if pending_space:
            yield from emit(" ", pending_space[0], pending_space[1])
            pending_space = None
        for item in transformed:
            yield from emit(item, offset, offset + 1)
            has_output = True


def _snippet(value: str, start: int, end: int, radius: int = 72) -> str:
    left = max(0, start - radius)
    right = min(len(value), end + radius)
    return ("…" if left else "") + value[left:right].replace("\n", " ").replace("\r", " ") + ("…" if right < len(value) else "")


def _source_id(base: str, match_index: int, start: int) -> str:
    return f"search:{base}:{match_index}:{start}"


def _locator_for_match(
    locator: dict[str, Any],
    source: str,
    match: TextMatch,
    file_type: str,
    *,
    line_start: int | None = None,
) -> dict[str, Any]:
    result = dict(locator)
    kind = str(result.get("kind") or file_type).casefold()
    raw_base = result.get("char_start")
    raw_base = raw_base if isinstance(raw_base, int) and not isinstance(raw_base, bool) else 0
    result["search_match"] = True
    result["search_range"] = {
        "coordinate_space": "source-text",
        "start": match.start,
        "end": match.end,
    }

    # These locators point into the same character stream as source_text.
    # Keep format-specific paragraph, node and table anchors intact elsewhere.
    if kind in {"txt", "md", "text", "pdf", "xml", "json", "html", "htm"}:
        base = result.get("char_start")
        base = base if isinstance(base, int) and not isinstance(base, bool) else 0
        result["char_start"] = base + match.start
        result["char_end"] = base + match.end
        if line_start is not None:
            result["line_start"] = line_start
            result["line_end"] = line_start + source[match.start:match.end].count("\n")

    if result.get("ocr") is True:
        absolute_start, absolute_end = raw_base + match.start, raw_base + match.end
        raw_map = result.get("ocr_map")
        if isinstance(raw_map, dict):
            word_boxes = raw_map.get("word_boxes")
            if isinstance(word_boxes, list):
                selected = [
                    box for box in word_boxes
                    if isinstance(box, list) and len(box) >= 8
                    and isinstance(box[4], int) and isinstance(box[5], int)
                    and box[4] < absolute_end and box[5] > absolute_start
                ]
                result["ocr_map"] = {**raw_map, "word_boxes": selected}
                lines = [box[6] for box in selected if isinstance(box[6], int)]
                if lines:
                    result["line_start"], result["line_end"] = min(lines), max(lines)
    return result


def search_source_blocks(
    sources: Iterable[Any],
    query: str,
    *,
    file_type: str,
    document_id: str,
    processing_version: int,
    offset: int = 0,
    limit: int = 50,
    excluded_pages: set[int] | None = None,
) -> dict[str, Any]:
    """Search the active source map, returning a page while counting every hit."""
    _validate_query(query, offset, limit)
    excluded_pages = excluded_pages or set()
    ordered = sorted(sources, key=lambda item: (int(getattr(item, "ordinal", 0)), str(getattr(item, "id", ""))))
    matches: list[dict[str, Any]] = []
    total = 0
    seen_sources: set[tuple[Any, ...]] = set()
    for source_index, source in enumerate(ordered):
        if bool(getattr(source, "is_derived", False)):
            continue
        raw_locator = getattr(source, "locator", None)
        locator = dict(raw_locator) if isinstance(raw_locator, dict) else {}
        if locator.get("page") in excluded_pages:
            continue
        text = locator.get("source_text")
        if not isinstance(text, str):
            text = getattr(source, "text", "")
        if not isinstance(text, str) or not text:
            continue
        source_identity = (
            tuple(locator.get(field) for field in ("page", "line_start", "char_start", "paragraph", "table", "row", "row_start", "sheet", "slide", "shape", "chapter", "path")),
            text,
        )
        if source_identity in seen_sources:
            continue
        seen_sources.add(source_identity)
        source_id = str(getattr(source, "id", f"source:{source_index}"))
        line_cursor = 0
        source_line = 1
        line_base = locator.get("line_start")
        line_base = line_base if isinstance(line_base, int) and not isinstance(line_base, bool) else None
        for local_index, match in enumerate(_iter_text_matches(text, query)):
            while line_cursor < match.start:
                if text[line_cursor] == "\n":
                    source_line += 1
                line_cursor += 1
            if offset <= total < offset + limit:
                match_line = line_base + source_line - 1 if line_base is not None else None
                mapped = _locator_for_match(locator, text, match, file_type, line_start=match_line)
                match_text = text[match.start:match.end]
                versioned = versioned_source_locator(
                    mapped,
                    document_id=document_id,
                    processing_version=processing_version,
                    file_type=file_type,
                )
                matches.append({
                    "id": _source_id(source_id, local_index, match.start),
                    "text": match_text,
                    "snippet": _snippet(text, match.start, match.end),
                    "locator": versioned,
                    "ordinal": int(getattr(source, "ordinal", 0)),
                    "is_derived": False,
                    "match_start": match.start,
                    "match_end": match.end,
                    "markdown_start": None,
                    "markdown_end": None,
                })
            total += 1
    return {"total": total, "offset": offset, "limit": limit, "matches": matches}


def search_markdown(
    markdown: str,
    query: str,
    *,
    chunks: Iterable[Any],
    document_id: str,
    processing_version: int,
    file_type: str,
    offset: int = 0,
    limit: int = 50,
) -> dict[str, Any]:
    """Search the complete Markdown artifact using its own global offsets."""
    _validate_query(query, offset, limit)
    ordered_chunks = sorted(chunks, key=lambda item: int(getattr(item, "ordinal", 0)))
    mapped_chunks = sorted(
        (
            (getattr(chunk, "markdown_char_start", None), getattr(chunk, "markdown_char_end", None), chunk)
            for chunk in ordered_chunks
            if isinstance(getattr(chunk, "markdown_char_start", None), int)
            and isinstance(getattr(chunk, "markdown_char_end", None), int)
        ),
        key=lambda item: item[0],
    )
    starts = [item[0] for item in mapped_chunks]
    matches: list[dict[str, Any]] = []
    total = 0
    line_cursor = 0
    line_number = 1
    for match in _iter_text_matches(markdown, query):
        while line_cursor < match.start:
            if markdown[line_cursor] == "\n":
                line_number += 1
            line_cursor += 1
        if offset <= total < offset + limit:
            locator: dict[str, Any] = {}
            source_ordinal = 0
            candidate_index = bisect_right(starts, match.start) - 1
            if candidate_index < 0 and mapped_chunks and mapped_chunks[0][0] < match.end:
                candidate_index = 0
            if candidate_index >= 0 and mapped_chunks[candidate_index][0] < match.end and mapped_chunks[candidate_index][1] > match.start:
                chunk = mapped_chunks[candidate_index][2]
                locator = dict(getattr(chunk, "locator", {}) or {})
                source_ordinal = int(getattr(chunk, "ordinal", 0))
            locator.update({
                "search_match": True,
                "markdown_char_start": match.start,
                "markdown_char_end": match.end,
                "markdown_line_start": line_number,
                "markdown_line_end": line_number + markdown[match.start:match.end].count("\n"),
                "match_range": {"coordinate_space": "markdown-codepoint", "start": match.start, "end": match.end},
            })
            versioned = versioned_source_locator(
                locator,
                document_id=document_id,
                processing_version=processing_version,
                file_type=file_type,
            )
            matches.append({
                "id": _source_id("markdown", total, match.start),
                "text": markdown[match.start:match.end],
                "snippet": _snippet(markdown, match.start, match.end),
                "locator": versioned,
                "ordinal": source_ordinal,
                "is_derived": False,
                "match_start": None,
                "match_end": None,
                "markdown_start": match.start,
                "markdown_end": match.end,
            })
        total += 1
    return {"total": total, "offset": offset, "limit": limit, "matches": matches}


def search_table_bytes(data: bytes, file_type: str, query: str, *, offset: int = 0, limit: int = 50) -> dict[str, Any]:
    """Search real spreadsheet cells and preserve sheet/row/column anchors."""
    _validate_query(query, offset, limit)
    matches: list[dict[str, Any]] = []
    total = 0

    def scan_rows(
        sheet_name: str | None,
        rows: Iterable[Iterable[Any]],
        *,
        omit_blank_rows: bool,
        preserve_physical_rows: bool = False,
    ) -> None:
        nonlocal total
        headers: list[str] | None = None
        visible_row = 0
        for physical_row, raw_row in enumerate(rows, start=1):
            row = ["" if value is None else str(value) for value in raw_row]
            is_empty = not any(cell.strip() for cell in row)
            if omit_blank_rows and is_empty:
                continue
            visible_row += 1
            row_number = physical_row if preserve_physical_rows or not omit_blank_rows else visible_row
            if headers is None:
                if omit_blank_rows and is_empty:
                    continue
                headers = [cell.strip() or f"Столбец {index + 1}" for index, cell in enumerate(row[:500])]
                if len(row) > 500:
                    raise DocumentParsingError("В таблице слишком много столбцов (максимум 500).")
            for column_index, cell in enumerate(row[:len(headers)]):
                for match in _iter_text_matches(cell, query):
                    column = headers[column_index]
                    locator = {
                        "kind": file_type,
                        "label": f"{column} · строка {row_number}",
                        "row": row_number,
                        "row_start": row_number,
                        "row_end": row_number,
                        "column": column,
                        "column_index": column_index,
                        "source_text": cell,
                        "search_match": True,
                        "search_range": {"coordinate_space": "table-cell-text", "start": match.start, "end": match.end},
                    }
                    if sheet_name is not None:
                        locator["sheet"] = sheet_name
                        locator["label"] = f"{sheet_name} · {column} · строка {row_number}"
                    if offset <= total < offset + limit:
                        matches.append({
                            "id": _source_id(f"{sheet_name or 'csv'}:{row_number}:{column_index}", 0, match.start),
                            "text": cell[match.start:match.end],
                            "snippet": _snippet(cell, match.start, match.end),
                            "locator": locator,
                            "ordinal": total,
                            "is_derived": False,
                            "match_start": match.start,
                            "match_end": match.end,
                            "markdown_start": None,
                            "markdown_end": None,
                        })
                    total += 1

    if file_type == "csv":
        from app.services.table_analysis import detect_csv_dialect
        text, _ = _decode_text_with_encoding(data)
        dialect, delimiter = detect_csv_dialect(text)
        try:
            scan_rows(
                None,
                csv.reader(io.StringIO(text, newline=""), dialect=dialect, delimiter=delimiter),
                omit_blank_rows=True,
                preserve_physical_rows=True,
            )
        except csv.Error as exc:
            raise DocumentParsingError("Не удалось разобрать строки CSV для поиска.") from exc
    elif file_type == "xlsx":
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise DocumentParsingError("Не удалось открыть XLSX для поиска.") from exc
        try:
            if len(workbook.worksheets) > 200:
                raise DocumentParsingError("В таблице слишком много листов (максимум 200).")
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) > 100_000 or (sheet.max_column or 0) > 500:
                    raise DocumentParsingError("Размер таблицы превышает безопасный предел строк или столбцов.")
                scan_rows(sheet.title, sheet.iter_rows(values_only=True), omit_blank_rows=False)
        finally:
            workbook.close()
    elif file_type == "xls":
        try:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        except Exception as exc:
            raise DocumentParsingError("Не удалось открыть XLS для поиска.") from exc
        try:
            if workbook.nsheets > 200:
                raise DocumentParsingError("В таблице слишком много листов (максимум 200).")
            for index in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(index)
                if sheet.nrows > 100_000 or sheet.ncols > 500:
                    raise DocumentParsingError("Размер таблицы превышает безопасный предел строк или столбцов.")
                scan_rows(sheet.name, (sheet.row_values(row) for row in range(sheet.nrows)), omit_blank_rows=False)
        finally:
            if hasattr(workbook, "release_resources"):
                workbook.release_resources()
    else:
        raise DocumentParsingError("Поиск ячеек доступен только для CSV, XLSX и XLS.")
    return {"total": total, "offset": offset, "limit": limit, "matches": matches}


def _validate_query(query: str, offset: int, limit: int) -> None:
    if not isinstance(query, str) or len(query) > MAX_SEARCH_QUERY_CHARS:
        raise DocumentParsingError("Поисковый запрос слишком длинный.")
    if not _normalise(query):
        raise DocumentParsingError("Введите текст для поиска.")
    if offset < 0:
        raise DocumentParsingError("offset не может быть отрицательным.")
    if limit < 1 or limit > MAX_SEARCH_PAGE_SIZE:
        raise DocumentParsingError("Размер страницы поиска должен быть от 1 до 100.")
