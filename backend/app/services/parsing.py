from __future__ import annotations

import bisect
import csv
import html
import io
import json
import re
import unicodedata
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from decimal import Decimal
from html.parser import HTMLParser
from pathlib import PurePosixPath
from typing import Any
from xml.etree.ElementTree import Element

from defusedxml import ElementTree as SafeElementTree
from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.services.numeric_values import parse_decimal

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".csv", ".xml",
    ".xlsx", ".xls", ".pptx", ".html", ".htm", ".json", ".epub",
}
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_DOCX_ENTRIES = 5_000
CHUNK_TARGET_CHARS = 1_100
CSV_GROUP_CHARS = 900
SPREADSHEET_MAX_ROWS = 100_000
SPREADSHEET_MAX_COLUMNS = 500
EPUB_MAX_ENTRIES = 5_000
EPUB_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024


class DocumentParsingError(ValueError):
    """An expected document validation or extraction failure."""


@dataclass(frozen=True)
class SourceBlock:
    text: str
    locator: dict[str, Any]
    derived: bool = False


@dataclass
class ParsedDocument:
    file_type: str
    blocks: list[SourceBlock] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)


def safe_filename(filename: str) -> str:
    leaf = filename.replace("\\", "/").split("/")[-1].strip().replace("\x00", "")
    leaf = re.sub(r"[\r\n\t]", " ", leaf).strip(" .")
    if not leaf:
        raise DocumentParsingError("Не удалось определить имя файла.")
    return leaf[:255]


def _decode_text_with_encoding(data: bytes) -> tuple[str, str]:
    if b"\x00" in data[:4096]:
        for encoding in ("utf-16", "utf-16-le", "utf-16-be"):
            try:
                return data.decode(encoding), encoding
            except UnicodeDecodeError:
                continue
        raise DocumentParsingError("Файл содержит неподдерживаемую двоичную кодировку.")
    encodings = ("utf-8-sig", "cp1251", "latin-1")
    for encoding in encodings:
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise DocumentParsingError("Не удалось определить кодировку текста.")


def _decode_text(data: bytes) -> str:
    return _decode_text_with_encoding(data)[0]


def _normalize_for_source_match(value: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", value).casefold().split())


def _find_normalized_source_range(
    source: str,
    expected: str,
    start_hint: int = 0,
    *,
    decode_entities: bool = False,
) -> tuple[int, int] | None:
    """Find a normalized citation in source text and return original offsets."""

    normalized: list[str] = []
    starts: list[int] = []
    ends: list[int] = []
    in_space = False
    def source_units() -> Iterable[tuple[str, int, int]]:
        offset = 0
        while offset < len(source):
            if decode_entities and source[offset] == "&":
                entity_end = source.find(";", offset + 1, min(len(source), offset + 20))
                if entity_end >= 0:
                    entity = source[offset:entity_end + 1]
                    decoded = html.unescape(entity)
                    if decoded != entity:
                        for character in decoded:
                            yield character, offset, entity_end + 1
                        offset = entity_end + 1
                        continue
            character = source[offset]
            yield character, offset, offset + 1
            offset += 1

    for character, offset, source_end in source_units():
        transformed = unicodedata.normalize("NFKC", character).casefold()
        if character.isspace() or transformed.isspace():
            if normalized and not in_space:
                normalized.append(" ")
                starts.append(offset)
                ends.append(source_end)
            elif normalized and in_space:
                ends[-1] = source_end
            in_space = True
            continue
        in_space = False
        for item in transformed:
            normalized.append(item)
            starts.append(offset)
            ends.append(source_end)
    needle = _normalize_for_source_match(expected)
    if not needle:
        return None
    source_normalized = "".join(normalized)
    candidates: list[int] = []
    offset = 0
    while (found := source_normalized.find(needle, offset)) >= 0:
        candidates.append(found)
        offset = found + 1
    if not candidates:
        return None
    candidate = min(candidates, key=lambda found: abs(starts[found] - start_hint))
    return starts[candidate], ends[candidate + len(needle) - 1]


def _split_long_text(text: str, locator: dict[str, Any], target: int = CHUNK_TARGET_CHARS) -> list[SourceBlock]:
    normalized = text.strip()
    if not normalized:
        return []
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", normalized) if part.strip()]
    if not paragraphs:
        paragraphs = [normalized]
    result: list[SourceBlock] = []
    buffer = ""
    part_index = 1
    base_char_start = locator.get("char_start")
    has_base_char_start = isinstance(base_char_start, int) and not isinstance(base_char_start, bool)
    leading_trim = len(text) - len(text.lstrip())

    def part_locator(start: int, end: int, part: int) -> dict[str, Any]:
        current = {**locator, "part": part}
        if has_base_char_start:
            current["char_start"] = base_char_start + leading_trim + start
            current["char_end"] = base_char_start + leading_trim + end
        else:
            current.pop("char_start", None)
            current.pop("char_end", None)
        return current

    for paragraph in paragraphs:
        if len(paragraph) > target:
            if buffer:
                result.append(SourceBlock(buffer, part_locator(0, len(buffer), part_index)))
                buffer = ""
                part_index += 1
            start = 0
            while start < len(paragraph):
                end = min(start + target, len(paragraph))
                if end < len(paragraph):
                    boundary = paragraph.rfind(" ", start + target // 2, end)
                    if boundary > start:
                        end = boundary
                part = paragraph[start:end].strip()
                if part:
                    result.append(SourceBlock(part, part_locator(start, end, part_index)))
                    part_index += 1
                start = max(end - 120, end)
            continue
        candidate = f"{buffer}\n\n{paragraph}" if buffer else paragraph
        if len(candidate) > target and buffer:
            result.append(SourceBlock(buffer, part_locator(0, len(buffer), part_index)))
            part_index += 1
            buffer = paragraph
        else:
            buffer = candidate
    if buffer:
        result.append(SourceBlock(buffer, part_locator(0, len(buffer), part_index)))
    if has_base_char_start:
        search_cursor = 0
        for block in result:
            source_range = _find_normalized_source_range(normalized, block.text, search_cursor)
            if source_range:
                start, end = source_range
                block.locator["char_start"] = base_char_start + leading_trim + start
                block.locator["char_end"] = base_char_start + leading_trim + end
                search_cursor = end
    return result


def _pdf_page_geometry(page: Any) -> dict[str, Any]:
    try:
        box = page.cropbox
        crop_box = [float(box.left), float(box.bottom), float(box.right), float(box.top)]
        width = abs(crop_box[2] - crop_box[0])
        height = abs(crop_box[3] - crop_box[1])
        rotation = int(getattr(page, "rotation", 0) or 0) % 360
    except (AttributeError, TypeError, ValueError, OverflowError):
        crop_box = None
        width = height = None
        rotation = 0
    if rotation in (90, 270):
        width, height = height, width
    return {"crop_box": crop_box, "rotation": rotation, "width": width, "height": height}


def _pdf_text_quality(text: str) -> tuple[bool, int, str | None]:
    visible = [character for character in text if not character.isspace()]
    alphanumeric = sum(character.isalnum() for character in visible)
    suspicious = sum(
        character == "\ufffd" or (ord(character) < 32 and character not in "\t\n\r")
        for character in visible
    )
    if not visible:
        return False, 0, "text_layer_missing"
    if alphanumeric < 24:
        return False, alphanumeric, "sparse_text_layer"
    if suspicious / len(visible) >= 0.12 or alphanumeric / len(visible) < 0.35:
        return False, alphanumeric, "suspicious_text_layer"
    return True, alphanumeric, None


def _parse_pdf(data: bytes) -> ParsedDocument:
    if not data.startswith(b"%PDF-"):
        raise DocumentParsingError("Расширение PDF не соответствует содержимому файла.")
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise DocumentParsingError("PDF защищён паролем. Зашифрованные документы пока не поддерживаются.")
        pages = list(reader.pages)
        from app.config import settings
        if len(pages) > settings.document_max_pages:
            raise DocumentParsingError('В PDF превышен безопасный предел страниц.')
        blocks: list[SourceBlock] = []
        extracted_chars = 0
        page_map: list[dict[str, Any]] = []
        ocr_pages: list[int] = []
        for page_number, page in enumerate(pages, start=1):
            has_content = page.get_contents() is not None
            page_text = "" if not has_content else page.extract_text(extraction_mode="layout") or page.extract_text() or ""
            extracted_chars += len(page_text.strip())
            if extracted_chars > settings.document_max_chars:
                raise DocumentParsingError('Текст PDF превышает безопасный предел.')
            native_quality, alphanumeric, reason = _pdf_text_quality(page_text)
            classification = "native" if native_quality else "ocr_candidate" if has_content else "blank"
            geometry = _pdf_page_geometry(page)
            page_map.append({
                "page": page_number,
                "classification": classification,
                "native_char_count": len(page_text.strip()),
                "native_alphanumeric_count": alphanumeric,
                "ocr_reason": reason if classification == "ocr_candidate" else None,
                **geometry,
            })
            if classification == "ocr_candidate":
                ocr_pages.append(page_number)
            blocks.extend(_split_long_text(page_text, {
                "kind": "pdf", "page": page_number, "label": f"Страница {page_number}",
                "char_start": 0, "char_end": len(page_text),
            }))
    except DocumentParsingError:
        raise
    except (PdfReadError, ValueError, OSError, KeyError) as exc:
        raise DocumentParsingError("PDF повреждён или имеет неподдерживаемую структуру.") from exc
    page_width = page_map[0].get("width") if page_map else None
    page_height = page_map[0].get("height") if page_map else None
    metadata: dict[str, Any] = {
        "page_count": len(pages),
        "pdf_page_map": page_map,
        "ocr_pages": ocr_pages,
        "ocr_required": bool(ocr_pages),
    }
    if page_width and page_height:
        metadata.update({"page_width": page_width, "page_height": page_height})
    if not blocks and ocr_pages:
        metadata["ocr_reason"] = "В PDF есть страницы без надёжного текстового слоя."
    elif not blocks:
        metadata["ocr_reason"] = "PDF содержит только пустые страницы."
    return ParsedDocument("pdf", blocks, metadata)


def _iter_docx_blocks(document: DocxDocumentType) -> Iterable[DocxParagraph | DocxTable]:
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield DocxParagraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield DocxTable(child, document)


def _validate_docx(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise DocumentParsingError("В DOCX слишком много внутренних частей.")
            if sum(item.file_size for item in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise DocumentParsingError("Внутренний размер DOCX превышает безопасный предел.")
            if any(item.flag_bits & 0x1 for item in entries):
                raise DocumentParsingError("Зашифрованные DOCX пока не поддерживаются.")
            names = {item.filename for item in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentParsingError("Файл не является документом DOCX.")
    except DocumentParsingError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise DocumentParsingError("DOCX повреждён или имеет неверную структуру.") from exc


def _parse_docx(data: bytes) -> ParsedDocument:
    _validate_docx(data)
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParsingError("Не удалось прочитать структуру DOCX.") from exc

    blocks: list[SourceBlock] = []
    paragraph_number = 0
    table_number = 0
    for item in _iter_docx_blocks(document):
        if isinstance(item, DocxParagraph):
            paragraph_number += 1
            text = item.text.strip()
            if text:
                blocks.extend(_split_long_text(text, {
                    "kind": "docx",
                    "label": f"Абзац {paragraph_number}",
                    "paragraph": paragraph_number,
                    "char_start": 0,
                    "char_end": len(text),
                    "heading": item.style.name if item.style and item.style.name.startswith("Heading") else None,
                }))
            continue
        table_number += 1
        for row_number, row in enumerate(item.rows, start=1):
            values = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
            row_text = " | ".join(value for value in values if value)
            if row_text:
                blocks.extend(_split_long_text(row_text, {
                    "kind": "docx_table",
                    "label": f"Таблица {table_number}, строка {row_number}",
                    "table": table_number,
                    "row": row_number,
                    "char_start": 0,
                    "char_end": len(row_text),
                }))
    if not blocks:
        raise DocumentParsingError("В DOCX не найден текст или заполненные строки таблиц.")
    metadata: dict[str, Any] = {"paragraph_count": paragraph_number, "table_count": table_number}
    if document.sections:
        section = document.sections[0]
        if section.page_width and section.page_height:
            metadata.update({
                "page_width": round(section.page_width.inches * 25.4, 2),
                "page_height": round(section.page_height.inches * 25.4, 2),
            })
    return ParsedDocument("docx", blocks, metadata)


def _parse_plain_text(data: bytes, file_type: str) -> ParsedDocument:
    text, encoding = _decode_text_with_encoding(data)
    raw_lines = text.splitlines(keepends=True)
    if not raw_lines and text:
        raw_lines = [text]
    lines = [line.rstrip("\r\n") for line in raw_lines]
    line_offsets: list[int] = []
    cursor = 0
    for raw_line in raw_lines:
        line_offsets.append(cursor)
        cursor += len(raw_line)
    blocks: list[SourceBlock] = []
    start_line = 1
    while start_line <= len(lines):
        end_line = min(start_line + 34, len(lines))
        section = "\n".join(lines[start_line - 1:end_line]).strip()
        if section:
            section_start = line_offsets[start_line - 1]
            section_end = section_start + len(section)
            section_blocks = _split_long_text(section, {
                "kind": file_type,
                "label": f"Строки {start_line}–{end_line}",
                "line_start": start_line,
                "line_end": end_line,
                "char_start": section_start,
                "char_end": section_end,
            })
            for block in section_blocks:
                match = _find_normalized_source_range(text, block.text, section_start)
                if match:
                    block.locator["char_start"], block.locator["char_end"] = match
                    block.locator["line_start"] = bisect.bisect_right(line_offsets, match[0])
                    block.locator["line_end"] = bisect.bisect_right(line_offsets, max(match[0], match[1] - 1))
            blocks.extend(section_blocks)
        start_line = end_line + 1
    if not blocks:
        raise DocumentParsingError("Файл пустой — извлекать нечего.")
    return ParsedDocument(file_type, blocks, {"line_count": len(lines), "encoding": encoding})


def _number(value: str) -> Decimal | None:
    return parse_decimal(value)


def _parse_csv(data: bytes) -> ParsedDocument:
    text, encoding = _decode_text_with_encoding(data)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        dialect = csv.excel
        delimiter = ";" if text.count(";") > text.count(",") else ","
    try:
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect=dialect, delimiter=delimiter))
    except csv.Error as exc:
        raise DocumentParsingError("Не удалось разобрать строки CSV.") from exc
    header_index = next((index for index, row in enumerate(rows) if any(cell.strip() for cell in row)), None)
    if header_index is None:
        raise DocumentParsingError("CSV пустой — строк не найдено.")

    header_row_number = header_index + 1
    headers = [cell.strip() or f"Столбец {index + 1}" for index, cell in enumerate(rows[header_index])]
    if len(headers) > 500:
        raise DocumentParsingError("В CSV слишком много столбцов (максимум 500).")
    data_rows = list(enumerate(rows[header_index + 1:], start=header_row_number + 1))
    blocks: list[SourceBlock] = []
    header_text = "Заголовки столбцов: " + " | ".join(headers)
    blocks.append(SourceBlock(header_text, {
        "kind": "csv",
        "label": "Заголовки столбцов",
        "row_start": header_row_number,
        "row_end": header_row_number,
        "columns": headers,
    }))

    group: list[str] = []
    group_start = header_row_number + 1
    group_chars = 0
    for row_number, row in data_rows:
        normalized = row[:len(headers)] + [""] * max(0, len(headers) - len(row))
        line = " | ".join(f"{headers[i]}: {normalized[i].strip()}" for i in range(len(headers)) if normalized[i].strip())
        if not line:
            continue
        if group and group_chars + len(line) > CSV_GROUP_CHARS:
            blocks.append(SourceBlock("\n".join(group), {
                "kind": "csv",
                "label": f"Строки {group_start}–{row_number - 1}",
                "row_start": group_start,
                "row_end": row_number - 1,
                "columns": headers,
            }))
            group = []
            group_chars = 0
            group_start = row_number
        if len(line) > CSV_GROUP_CHARS:
            if group:
                blocks.append(SourceBlock("\n".join(group), {
                    "kind": "csv",
                    "label": f"Строки {group_start}–{row_number - 1}",
                    "row_start": group_start,
                    "row_end": row_number - 1,
                    "columns": headers,
                }))
                group = []
                group_chars = 0
            blocks.extend(_split_long_text(line, {
                "kind": "csv",
                "label": f"Строка {row_number}",
                "row_start": row_number,
                "row_end": row_number,
                "columns": headers,
            }, target=CSV_GROUP_CHARS))
            group_start = row_number + 1
            continue
        group.append(line)
        group_chars += len(line) + 1
    if group:
        blocks.append(SourceBlock("\n".join(group), {
            "kind": "csv",
            "label": f"Строки {group_start}–{header_row_number + len(data_rows)}",
            "row_start": group_start,
            "row_end": header_row_number + len(data_rows),
            "columns": headers,
        }))

    numeric_columns: list[dict[str, Any]] = []
    for column_index, header in enumerate(headers):
        values = [_number(row[column_index]) for _row_number, row in data_rows if column_index < len(row)]
        numbers = [value for value in values if value is not None]
        nonempty_count = sum(1 for _row_number, row in data_rows if column_index < len(row) and row[column_index].strip())
        if numbers and len(numbers) >= max(2, int(nonempty_count * 0.75)):
            total = sum(numbers, Decimal(0))
            numeric_columns.append({
                "name": header,
                "count": len(numbers),
                "sum": str(total),
                "average": str(total / Decimal(len(numbers))),
                "minimum": str(min(numbers)),
                "maximum": str(max(numbers)),
            })

    metadata = {
        "row_count": len(data_rows),
        "column_count": len(headers),
        "columns": headers,
        "header_row": header_row_number,
        "data_start_row": header_row_number + 1,
        "row_end": header_row_number + len(data_rows),
        "delimiter": delimiter,
        "encoding": encoding,
        "numeric_columns": numeric_columns,
    }
    return ParsedDocument("csv", blocks, metadata)


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _parse_xml(data: bytes) -> ParsedDocument:
    try:
        root = SafeElementTree.fromstring(data)
    except Exception as exc:
        raise DocumentParsingError("XML повреждён или содержит запрещённую DTD/entity-конструкцию.") from exc

    source_text, encoding = _decode_text_with_encoding(data)
    line_starts = [0]
    line_starts.extend(index + 1 for index, char in enumerate(source_text) if char == "\n")
    source_cursor = 0
    blocks: list[SourceBlock] = []

    def value_locator(path: str, value: str) -> dict[str, Any]:
        nonlocal source_cursor
        locator: dict[str, Any] = {"kind": "xml", "label": path, "path": path}
        source_range = _find_normalized_source_range(source_text, value, source_cursor, decode_entities=True)
        if source_range:
            start, end = source_range
            locator.update({
                "char_start": start,
                "char_end": end,
                "line_start": bisect.bisect_right(line_starts, start),
                "line_end": bisect.bisect_right(line_starts, max(start, end - 1)),
            })
            source_cursor = end
        return locator
    stack: list[tuple[Element, str]] = [(root, f"/{_local_name(root.tag)}[1]")]
    node_count = 0
    while stack:
        element, path = stack.pop()
        node_count += 1
        if node_count > 200_000:
            raise DocumentParsingError("В XML слишком много узлов (максимум 200 000).")
        children = list(element)
        attributes = "; ".join(f"{_local_name(name)}={value}" for name, value in element.attrib.items())
        text = (element.text or "").strip()
        if text or (attributes and not children):
            value = text or attributes
            if attributes and text:
                value = f"{text} ({attributes})"
            # Attribute summaries do not occur literally in the source XML;
            # locate the visible element text and retain a path locator for
            # attribute-only values.
            locator = value_locator(path, text) if text else {"kind": "xml", "label": path, "path": path}
            if attributes and text:
                # The indexed summary includes attributes that are not at this
                # text offset in the original XML, so only the path/line is a
                # reliable locator for this combined block.
                locator.pop("char_start", None)
                locator.pop("char_end", None)
                locator["match_quality"] = "approximate"
            xml_blocks = _split_long_text(value, locator)
            if not attributes and text and isinstance(locator.get("char_start"), int):
                source_hint = locator["char_start"]
                for block in xml_blocks:
                    source_range = _find_normalized_source_range(
                        source_text, block.text, source_hint, decode_entities=True,
                    )
                    if source_range:
                        start, end = source_range
                        block.locator.update({
                            "char_start": start,
                            "char_end": end,
                            "line_start": bisect.bisect_right(line_starts, start),
                            "line_end": bisect.bisect_right(line_starts, max(start, end - 1)),
                        })
                        source_hint = end
            blocks.extend(xml_blocks)
        if element.tail and element.tail.strip():
            tail = element.tail.strip()
            tail_locator = value_locator(path, tail)
            tail_blocks = _split_long_text(tail, tail_locator)
            if isinstance(tail_locator.get("char_start"), int):
                source_hint = tail_locator["char_start"]
                for block in tail_blocks:
                    source_range = _find_normalized_source_range(
                        source_text, block.text, source_hint, decode_entities=True,
                    )
                    if source_range:
                        start, end = source_range
                        block.locator.update({
                            "char_start": start,
                            "char_end": end,
                            "line_start": bisect.bisect_right(line_starts, start),
                            "line_end": bisect.bisect_right(line_starts, max(start, end - 1)),
                        })
                        source_hint = end
            blocks.extend(tail_blocks)
        counts: dict[str, int] = {}
        indexed_children: list[tuple[Element, str]] = []
        for child in children:
            name = _local_name(child.tag)
            counts[name] = counts.get(name, 0) + 1
            indexed_children.append((child, f"{path}/{name}[{counts[name]}]"))
        stack.extend(reversed(indexed_children))
    if not blocks:
        raise DocumentParsingError("В XML не найдено текстовых значений или атрибутов.")
    metadata: dict[str, Any] = {"root": _local_name(root.tag), "element_count": node_count}
    metadata["encoding"] = encoding
    return ParsedDocument("xml", blocks, metadata)


def _parse_json(data: bytes) -> ParsedDocument:
    text, encoding = _decode_text_with_encoding(data)
    try:
        value = json.loads(text)
    except json.JSONDecodeError as exc:
        raise DocumentParsingError("JSON повреждён или имеет неверный синтаксис.") from exc
    pretty = json.dumps(value, ensure_ascii=False, indent=2)
    if not pretty.strip():
        raise DocumentParsingError("JSON пустой — извлекать нечего.")
    blocks: list[SourceBlock] = []

    line_starts = [0]
    line_starts.extend(index + 1 for index, char in enumerate(text) if char == "\n")

    def line_at(offset: int) -> int:
        return bisect.bisect_right(line_starts, offset)

    decoder = json.JSONDecoder()

    def skip_space(offset: int) -> int:
        while offset < len(text) and text[offset].isspace():
            offset += 1
        return offset

    def walk(path: str, offset: int) -> int:
        offset = skip_space(offset)
        if offset >= len(text):
            return offset
        token = text[offset]
        if token == "{":
            offset = skip_space(offset + 1)
            while offset < len(text) and text[offset] != "}":
                key, key_end = decoder.raw_decode(text, offset)
                offset = skip_space(key_end)
                if offset >= len(text) or text[offset] != ":":
                    raise ValueError("Invalid object token")
                offset = walk(f"{path}.{key}", offset + 1)
                offset = skip_space(offset)
                if offset < len(text) and text[offset] == ",":
                    offset = skip_space(offset + 1)
                else:
                    break
            if offset < len(text) and text[offset] == "}":
                return offset + 1
            raise ValueError("Unclosed object token")
        if token == "[":
            offset = skip_space(offset + 1)
            index = 0
            while offset < len(text) and text[offset] != "]":
                offset = walk(f"{path}[{index}]", offset)
                index += 1
                offset = skip_space(offset)
                if offset < len(text) and text[offset] == ",":
                    offset = skip_space(offset + 1)
                else:
                    break
            if offset < len(text) and text[offset] == "]":
                return offset + 1
            raise ValueError("Unclosed array token")

        start = offset
        scalar, end = decoder.raw_decode(text, offset)
        value_text = json.dumps(scalar, ensure_ascii=False) if scalar is not None else "null"
        blocks.append(SourceBlock(str(value_text), {
            "kind": "json",
            "label": path,
            "path": path,
            "line_start": line_at(start),
            "line_end": line_at(max(start, end - 1)),
            "char_start": start,
            "char_end": end,
        }))
        return end

    try:
        walk("$", 0)
    except (ValueError, json.JSONDecodeError):
        # json.loads already validated the document. Keep path-based locators
        # as a safe compatibility fallback if tokenizer behavior ever differs.
        blocks.clear()

        def walk_fallback(node: Any, path: str) -> None:
            if isinstance(node, dict):
                for key, child in node.items():
                    walk_fallback(child, f"{path}.{key}")
            elif isinstance(node, list):
                for index, child in enumerate(node):
                    walk_fallback(child, f"{path}[{index}]")
            else:
                value_text = json.dumps(node, ensure_ascii=False) if node is not None else "null"
                blocks.append(SourceBlock(str(value_text), {"kind": "json", "label": path, "path": path}))

        walk_fallback(value, "$")
    if not blocks:
        blocks.append(SourceBlock(pretty, {"kind": "json", "label": "JSON-документ", "path": "$"}))
    return ParsedDocument("json", blocks, {
        "encoding": encoding,
        "line_count": len(pretty.splitlines()),
        "root_type": type(value).__name__,
    })


class _HtmlBlockParser(HTMLParser):
    _BLOCK_TAGS: frozenset[str] = frozenset({
        "address", "article", "aside", "blockquote", "br", "dd", "div", "dl", "dt",
        "figcaption", "figure", "footer", "h1", "h2", "h3", "h4", "h5", "h6",
        "header", "hr", "li", "main", "nav", "ol", "p", "pre", "section", "table",
        "td", "th", "tr", "ul",
    })

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple[str, str, int, int, int]] = []
        self._parts: list[tuple[str, int, int]] = []
        self._tag = "body"
        self._start_line = 1
        self._ignored: list[str] = []
        self.raw_text = ""
        self.line_starts = [0]

    def set_source(self, text: str) -> None:
        self.raw_text = text
        self.line_starts = [0]
        self.line_starts.extend(index + 1 for index, char in enumerate(text) if char == "\n")

    def _flush(self) -> None:
        text = re.sub(r"\s+", " ", " ".join(part[0] for part in self._parts)).strip()
        if text:
            start = min(part[1] for part in self._parts)
            end = max(part[2] for part in self._parts)
            self.blocks.append((text, self._tag, self._start_line, start, end))
        self._parts = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {'script', 'style', 'iframe', 'object', 'noscript'}:
            self._ignored.append(tag.lower())
        if self._ignored:
            return
        if tag.lower() in self._BLOCK_TAGS and self._parts:
            self._flush()
        if tag.lower() in self._BLOCK_TAGS:
            self._tag = tag.lower()
            self._start_line = self.getpos()[0]

    def handle_endtag(self, tag: str) -> None:
        if self._ignored:
            if tag.lower() == self._ignored[-1]:
                self._ignored.pop()
            return
        if tag.lower() in self._BLOCK_TAGS:
            self._flush()

    def handle_data(self, data: str) -> None:
        if data.strip() and not self._ignored:
            line, column = self.getpos()
            line_index = min(max(line - 1, 0), len(self.line_starts) - 1)
            start = self.line_starts[line_index] + column
            end = self.raw_text.find("<", start)
            if end < 0:
                end = len(self.raw_text)
            self._parts.append((data, start, end))

    def close(self) -> None:
        super().close()
        self._flush()


def _parse_html(data: bytes, file_type: str = "html") -> ParsedDocument:
    text, encoding = _decode_text_with_encoding(data)
    parser = _HtmlBlockParser()
    parser.set_source(text)
    try:
        parser.feed(text)
        parser.close()
    except Exception as exc:
        raise DocumentParsingError("HTML повреждён или имеет неподдерживаемую структуру.") from exc
    blocks = [SourceBlock(value, {
        "kind": "html",
        "label": f"{tag.upper()} · строка {line}",
        "element": tag,
        "line_start": line,
        "line_end": bisect.bisect_right(parser.line_starts, max(start, end - 1)),
        "char_start": start,
        "char_end": end,
    }) for value, tag, line, start, end in parser.blocks]
    if not blocks:
        raise DocumentParsingError("В HTML не найден текст.")
    return ParsedDocument(file_type, blocks, {
        "encoding": encoding,
        "line_count": len(text.splitlines()),
    })


def _spreadsheet_blocks(
    rows_by_sheet: Iterable[tuple[str, list[list[str]]]],
    file_type: str,
    formula_cells_by_sheet: dict[str, dict[int, list[dict[str, Any]]]] | None = None,
) -> ParsedDocument:
    blocks: list[SourceBlock] = []
    sheet_count = 0
    total_rows = 0
    max_columns = 0
    formula_count = 0
    formula_cache_missing_count = 0
    columns_by_sheet: dict[str, list[str]] = {}
    numeric_columns: list[dict[str, Any]] = []
    if formula_cells_by_sheet is None:
        formula_cells_by_sheet = {}
    for sheet_name, rows in rows_by_sheet:
        sheet_count += 1
        if sheet_count > 200:
            raise DocumentParsingError("В таблице слишком много листов (максимум 200).")
        if not rows or not any(str(cell).strip() for cell in rows[0]):
            continue
        if len(rows) > SPREADSHEET_MAX_ROWS:
            raise DocumentParsingError("В таблице слишком много строк (максимум 100 000 на лист).")
        headers = [str(cell).strip() or f"Столбец {index + 1}" for index, cell in enumerate(rows[0])]
        headers = headers[:SPREADSHEET_MAX_COLUMNS]
        columns_by_sheet[sheet_name] = headers
        max_columns = max(max_columns, len(headers))
        data_rows = rows[1:]
        total_rows += len(data_rows)
        if not numeric_columns:
            for column_index, header in enumerate(headers):
                values = [_number(row[column_index]) for row in data_rows if column_index < len(row)]
                numbers = [value for value in values if value is not None]
                if numbers and len(numbers) >= max(2, int(max(1, len(data_rows)) * 0.75)):
                    total = sum(numbers, Decimal(0))
                    numeric_columns.append({
                        "name": header,
                        "count": len(numbers),
                        "sum": str(total),
                        "average": str(total / Decimal(len(numbers))),
                        "minimum": str(min(numbers)),
                        "maximum": str(max(numbers)),
                    })
        header_formulas = formula_cells_by_sheet.get(sheet_name, {}).get(1, [])
        header_locator: dict[str, Any] = {
            "kind": file_type,
            "label": f"Лист «{sheet_name}»",
            "sheet": sheet_name,
            "row": 1,
            "row_start": 1,
            "row_end": 1,
        }
        if file_type == "xlsx":
            header_locator["sheet_row_count"] = len(rows)
            header_locator["populated_columns"] = [
                index for index, value in enumerate(rows[0][:len(headers)])
                if str(value).strip()
            ]
        if header_formulas:
            header_locator["formula_cells"] = [
                {
                    **item,
                    "column": headers[item["column_index"]] if item["column_index"] < len(headers) else f"Столбец {item['column_index'] + 1}",
                }
                for item in header_formulas
            ]
            formula_count += len(header_formulas)
            formula_cache_missing_count += sum(not item["has_cached_value"] for item in header_formulas)
        blocks.append(SourceBlock(f"Лист «{sheet_name}»: " + " | ".join(headers), header_locator))
        for row_number, row in enumerate(data_rows, start=2):
            normalized = [str(value).strip() for value in row[:len(headers)]]
            normalized.extend([""] * max(0, len(headers) - len(normalized)))
            formula_cells = [
                {
                    **item,
                    "column": headers[item["column_index"]] if item["column_index"] < len(headers) else f"Столбец {item['column_index'] + 1}",
                }
                for item in formula_cells_by_sheet.get(sheet_name, {}).get(row_number, [])
            ]
            formulas_by_column = {item["column_index"]: item for item in formula_cells}
            values: list[str] = []
            for index, value in enumerate(normalized):
                formula = formulas_by_column.get(index)
                if formula is None:
                    display = value
                elif formula["has_cached_value"]:
                    cached_value = formula["cached_value"]
                    cached = str(cached_value) if cached_value is not None else "пустое значение"
                    display = f"{value} (формула: {formula['formula']}; сохранённое значение: {cached})"
                else:
                    display = (
                        f"формула: {formula['formula']}; сохранённое значение отсутствует "
                        "(приложение формулу не вычисляло)"
                    )
                if display:
                    values.append(f"{headers[index]}: {display}")
            line = " | ".join(values)
            formula_count += len(formula_cells)
            formula_cache_missing_count += sum(not item["has_cached_value"] for item in formula_cells)
            if line:
                locator: dict[str, Any] = {
                    "kind": file_type,
                    "label": f"Лист «{sheet_name}», строка {row_number}",
                    "sheet": sheet_name,
                    "row": row_number,
                    "row_start": row_number,
                    "row_end": row_number,
                    "columns": headers,
                }
                if file_type == "xlsx":
                    locator["populated_columns"] = [
                        index for index, value in enumerate(normalized)
                        if value or index in formulas_by_column
                    ]
                if formula_cells:
                    locator["formula_cells"] = formula_cells
                blocks.append(SourceBlock(line, locator))
    if not blocks:
        raise DocumentParsingError("В таблице не найдено заполненных листов.")
    return ParsedDocument(file_type, blocks, {
        "sheet_count": sheet_count,
        "row_count": total_rows,
        "column_count": max_columns,
        "columns_by_sheet": columns_by_sheet,
        "columns": next(iter(columns_by_sheet.values()), []),
        "numeric_columns": numeric_columns,
        "formula_count": formula_count,
        "formula_cache_missing_count": formula_cache_missing_count,
    })


def _parse_xlsx(data: bytes) -> ParsedDocument:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DocumentParsingError("Для XLSX не установлен модуль openpyxl.") from exc
    formulas_workbook = None
    values_workbook = None
    try:
        formulas_workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        values_workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        if len(formulas_workbook.worksheets) > 200 or any(
            (sheet.max_row or 0) > SPREADSHEET_MAX_ROWS or (sheet.max_column or 0) > SPREADSHEET_MAX_COLUMNS
            for sheet in formulas_workbook.worksheets
        ):
            raise DocumentParsingError('Размер таблицы превышает безопасный предел строк, столбцов или листов.')
        formula_cells_by_sheet: dict[str, dict[int, list[dict[str, Any]]]] = {}

        def sheet_rows() -> Iterable[tuple[str, list[list[str]]]]:
            for formula_sheet in formulas_workbook.worksheets:
                value_sheet = values_workbook[formula_sheet.title]
                rows: list[list[str]] = []
                formula_rows: dict[int, list[dict[str, Any]]] = {}
                formula_iterator = formula_sheet.iter_rows(values_only=False)
                value_iterator = value_sheet.iter_rows(values_only=False)
                for row_number, (formula_row, value_row) in enumerate(
                    zip(formula_iterator, value_iterator, strict=True),
                    start=1,
                ):
                    row_values: list[str] = []
                    for column_index, (formula_cell, value_cell) in enumerate(zip(formula_row, value_row, strict=True)):
                        cached_value = value_cell.value
                        row_values.append("" if cached_value is None else str(cached_value))
                        if formula_cell.data_type != "f":
                            continue
                        raw_formula = formula_cell.value
                        formula = str(getattr(raw_formula, "text", raw_formula))
                        if not formula.startswith("="):
                            formula = f"={formula}"
                        formula_rows.setdefault(row_number, []).append({
                            "cell": formula_cell.coordinate,
                            "column_index": column_index,
                            "formula": formula,
                            "has_cached_value": cached_value is not None,
                            "cached_value": None if cached_value is None else str(cached_value),
                        })
                    rows.append(row_values)
                formula_cells_by_sheet[formula_sheet.title] = formula_rows
                yield formula_sheet.title, rows

        return _spreadsheet_blocks(sheet_rows(), "xlsx", formula_cells_by_sheet)
    except DocumentParsingError:
        raise
    except Exception as exc:
        raise DocumentParsingError("XLSX повреждён или имеет неверную структуру.") from exc
    finally:
        if formulas_workbook is not None:
            formulas_workbook.close()
        if values_workbook is not None:
            values_workbook.close()


def _parse_xls(data: bytes) -> ParsedDocument:
    try:
        import xlrd
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    except ImportError as exc:
        raise DocumentParsingError("Для XLS не установлен модуль xlrd.") from exc
    except Exception as exc:
        raise DocumentParsingError("XLS повреждён или имеет неверную структуру.") from exc
    rows = []
    for sheet in workbook.sheets():
        if sheet.nrows > SPREADSHEET_MAX_ROWS or sheet.ncols > SPREADSHEET_MAX_COLUMNS or workbook.nsheets > 200:
            raise DocumentParsingError('Размер таблицы превышает безопасный предел строк, столбцов или листов.')
        rows.append((sheet.name, [["" if value is None else str(value) for value in sheet.row_values(index)] for index in range(sheet.nrows)]))
    return _spreadsheet_blocks(rows, "xls")


def _parse_pptx(data: bytes) -> ParsedDocument:
    from app.config import settings
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise DocumentParsingError("Для PPTX не установлен модуль python-pptx.") from exc
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParsingError("PPTX повреждён или имеет неверную структуру.") from exc
    blocks: list[SourceBlock] = []
    if len(presentation.slides) > settings.document_max_pages:
        raise DocumentParsingError('В презентации слишком много слайдов.')
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape_number, shape in enumerate(slide.shapes, start=1):
            text = ""
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
            elif getattr(shape, "has_table", False):
                text = " | ".join(cell.text.strip() for row in shape.table.rows for cell in row.cells if cell.text.strip())
            if text:
                blocks.append(SourceBlock(text, {
                    "kind": "pptx",
                    "label": f"Слайд {slide_number}, блок {shape_number}",
                    "slide": slide_number,
                    "shape": shape_number,
                }))
    if not blocks:
        raise DocumentParsingError("В PPTX не найден текст.")
    return ParsedDocument("pptx", blocks, {
        "slide_count": len(presentation.slides),
        "page_width": round(presentation.slide_width / 914400 * 25.4, 2),
        "page_height": round(presentation.slide_height / 914400 * 25.4, 2),
    })


def _safe_epub_member(name: str) -> str:
    path = PurePosixPath(name)
    if path.is_absolute() or ".." in path.parts:
        raise DocumentParsingError("EPUB содержит небезопасный путь.")
    return str(path)


def _parse_epub(data: bytes) -> ParsedDocument:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise DocumentParsingError("EPUB повреждён или имеет неверную структуру.") from exc
    with archive:
        entries = archive.infolist()
        if len(entries) > EPUB_MAX_ENTRIES or sum(item.file_size for item in entries) > EPUB_MAX_UNCOMPRESSED_BYTES:
            raise DocumentParsingError("Размер внутреннего содержимого EPUB превышает безопасный предел.")
        html_entries = [item for item in entries if item.filename.lower().endswith((".xhtml", ".html", ".htm"))]
        if not html_entries:
            raise DocumentParsingError("В EPUB не найдено содержимое глав.")
        blocks: list[SourceBlock] = []
        for chapter, item in enumerate(html_entries, start=1):
            member = _safe_epub_member(item.filename)
            chapter_data = archive.read(item.filename)
            chapter_text, _ = _decode_text_with_encoding(chapter_data)
            line_starts = [0]
            line_starts.extend(index + 1 for index, char in enumerate(chapter_text) if char == "\n")
            for value, tag, line, start, end in _html_parser_blocks(chapter_data):
                blocks.append(SourceBlock(value, {
                    "kind": "epub",
                    "label": f"Глава {chapter} · {tag.upper()} · строка {line}",
                    "chapter": chapter,
                    "path": member,
                    "element": tag,
                    "line_start": line,
                    "line_end": bisect.bisect_right(line_starts, max(start, end - 1)),
                    "char_start": start,
                    "char_end": end,
                }))
    if not blocks:
        raise DocumentParsingError("В EPUB не найден текст глав.")
    return ParsedDocument("epub", blocks, {"chapter_count": len(html_entries)})


def _html_parser_blocks(data: bytes) -> list[tuple[str, str, int, int, int]]:
    text = _decode_text(data)
    parser = _HtmlBlockParser()
    parser.set_source(text)
    parser.feed(text)
    parser.close()
    return parser.blocks


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    filename = safe_filename(filename)
    extension = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentParsingError(f"Формат {extension or 'без расширения'} не поддерживается. Допустимы: {supported}.")
    if not data:
        raise DocumentParsingError("Файл пустой.")
    from app.services.document_security import validate_content
    validate_content(filename, data)
    if extension == ".pdf":
        return _parse_pdf(data)
    if extension == ".docx":
        return _parse_docx(data)
    if extension == ".txt":
        return _parse_plain_text(data, "txt")
    if extension == ".md":
        return _parse_plain_text(data, "md")
    if extension == ".csv":
        return _parse_csv(data)
    if extension == ".xml":
        return _parse_xml(data)
    if extension == ".json":
        return _parse_json(data)
    if extension in {".html", ".htm"}:
        return _parse_html(data, extension.removeprefix("."))
    if extension == ".xlsx":
        return _parse_xlsx(data)
    if extension == ".xls":
        return _parse_xls(data)
    if extension == ".pptx":
        return _parse_pptx(data)
    return _parse_epub(data)
